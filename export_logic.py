# -*- coding: utf-8 -*-
"""
Logic xuat file Excel (.xlsx) va Word (.docx) cho bao gia.
Dung lai theo dung layout mau: trang bao gia chinh, cac trang du toan
chi tiet theo hang muc (bang Tong hop chi phi + Dien giai chi phi chuyen
gia + Dien giai chi phi khac), va trang Bang luong chuyen gia dung chung.
"""

import math
import base64
import unicodedata
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _decode_logo_bytes(data_uri):
    """Giai ma logo tu data URI (vd 'data:image/png;base64,....') thanh bytes.
    Tra ve None neu khong co / khong hop le, de cac noi goi co the fallback an toan."""
    if not data_uri or not isinstance(data_uri, str) or "," not in data_uri:
        return None
    try:
        _, b64data = data_uri.split(",", 1)
        return base64.b64decode(b64data)
    except Exception:
        return None


def _resolve_footer_caption(state):
    """Dong bo voi JS: chu chan trang uu tien lay theo footerCaption nguoi dung nhap,
    neu de trong thi lay theo Ten cong ty (viet hoa)."""
    caption = (state.get("footerCaption") or "").strip()
    return caption or (state.get("companyName", "") or "").upper()


def fmt_number(n):
    try:
        return f"{round(n):,}".replace(",", ".")
    except (TypeError, ValueError):
        return str(n)


def compute_row_numbers(rows):
    """Dong bo voi JS computeRowNumbers: danh so 1, 1.1, 1.2, 2, 2.1...
    (CHI danh cho otherRows - toi da 2 cap qua co "isSub". expertRows dung
    compute_expert_row_numbers() nhieu cap o duoi.)"""
    main_num = 0
    sub_num = 0
    numbers = []
    for r in rows:
        if r.get("isSub") and main_num > 0:
            sub_num += 1
            numbers.append(f"{main_num}.{sub_num}")
        else:
            main_num += 1
            sub_num = 0
            numbers.append(f"{main_num}")
    return numbers


def compute_expert_row_numbers(rows):
    """Dong bo voi JS computeExpertRowNumbers: danh so NHIEU CAP cho expertRows, dua tren
    truong 'depth' (0 = khong thut, 1 = con, 2 = chau...), vd: 1, 1.1, 1.1.1, 1.1.2, 1.2, 2..."""
    counters = []
    numbers = []
    for r in rows:
        d = int(r.get("depth") or 0)
        counters = counters[: d + 1] + [0] * (d + 1 - len(counters)) if d + 1 > len(counters) else counters[: d + 1]
        counters[d] = counters[d] + 1 if len(counters) > d and counters[d] else 1
        numbers.append(".".join(str(c) for c in counters))
    return numbers


def row_own_amount(row, daily_rate):
    """Thanh tien 1 dong cong viec = tong theo TAT CA muc chuyen gia trong dong do.
    Ho tro ca cau truc moi (levels[]) va cau truc cu (levelId/quantity/days phang)."""
    levels = row.get("levels")
    if levels:
        return sum(daily_rate(lv.get("levelId")) * (lv.get("quantity") or 0) * (lv.get("days") or 0) for lv in levels)
    return daily_rate(row.get("levelId")) * (row.get("quantity") or 0) * (row.get("days") or 0)


def analyze_expert_rows(expert_rows, daily_rate):
    """Dong bo voi JS analyzeExpertRows: phan tich cau truc NHIEU CAP (khong gioi han o 2 cap
    me/con nhu truoc). 1 dong co cac dong CON TRUC TIEP (depth = depth cua no + 1, nam ngay sau
    no lien tiep cho den khi gap 1 dong co depth <= depth cua no) thi KHONG cong gia tri rieng
    cua no, ma hien TONG cua CAC DONG CON TRUC TIEP do (de quy qua displayAmount cua tung con -
    neu 1 dong con lai co chau rieng, tong cua no da tu cong gop chau dung roi)."""
    n = len(expert_rows)
    depths = [int(r.get("depth") or 0) for r in expert_rows]
    analysis = [None] * n

    def compute(idx):
        if analysis[idx] is not None:
            return analysis[idx]
        r = expert_rows[idx]
        d = depths[idx]
        child_indices = []
        j = idx + 1
        while j < n and depths[j] > d:
            if depths[j] == d + 1:
                child_indices.append(j)
            j += 1
        has_children = bool(child_indices)
        children_total = 0
        for ci in child_indices:
            children_total += compute(ci)["displayAmount"]
        own_amount = row_own_amount(r, daily_rate)
        display_amount = children_total if has_children else own_amount
        contributes = 0 if has_children else own_amount
        analysis[idx] = {
            "hasChildren": has_children, "childrenTotal": children_total,
            "ownAmount": own_amount, "displayAmount": display_amount, "contributes": contributes,
            "depth": d, "childIndices": child_indices,
        }
        return analysis[idx]

    for idx in range(n):
        compute(idx)
    return analysis


# Bang tham chieu ty le % Chi phi quan ly (Cql) theo Chi phi chuyen gia (Ccg) mac dinh -
# dung khi state khong co "cqlTiers" (file cu). Phai khop CHINH XAC voi mac dinh cqlTiers
# trong ui.html.
DEFAULT_CQL_TIERS = [
    {"upTo": 1000000000, "pct": 55},
    {"upTo": 5000000000, "pct": 50},
    {"upTo": None, "pct": 45},
]


def resolve_cql_pct(ccg, cql_tiers):
    """Tu dong chon ty le % Cql theo Ccg, dua tren bang tham chieu cql_tiers - dong bo
    tuyet doi voi resolveCqlPct() trong ui.html (xem giai thich chi tiet o do)."""
    tiers = list(cql_tiers) if cql_tiers else list(DEFAULT_CQL_TIERS)
    if not tiers:
        return 0

    def sort_key(t):
        return (t.get("upTo") is None, t.get("upTo") if t.get("upTo") is not None else 0)

    sorted_tiers = sorted(tiers, key=sort_key)
    n = len(sorted_tiers)
    for i, tier in enumerate(sorted_tiers):
        is_last = i == n - 1
        up_to = tier.get("upTo")
        if is_last or up_to is None:
            return tier.get("pct", 0)
        if i == 0:
            if ccg <= up_to:
                return tier.get("pct", 0)
        else:
            prev_up_to = sorted_tiers[i - 1].get("upTo")
            if ccg > prev_up_to and ccg < up_to:
                return tier.get("pct", 0)
    return sorted_tiers[-1].get("pct", 0)


def _cql_pct_formula(ccg_ref, cql_tiers):
    """Dung 1 cong thuc Excel IF long nhau tra ve % Cql THEO DUNG Ccg cua o tham chieu ccg_ref (vd 'E4')
    - dong bo tuyet doi voi resolve_cql_pct()/resolveCqlPct() o tren, nhung la CONG THUC SONG: neu nguoi
    dung sua so lieu lam Ccg thay doi qua nguong bac thang, % Cql trong file Excel se TU DONG doi theo,
    khong con la con so co dinh ghi san."""
    tiers = list(cql_tiers) if cql_tiers else list(DEFAULT_CQL_TIERS)
    if not tiers:
        return "0"

    def sort_key(t):
        return (t.get("upTo") is None, t.get("upTo") if t.get("upTo") is not None else 0)

    tiers = sorted(tiers, key=sort_key)
    n = len(tiers)
    expr = str(tiers[-1].get("pct", 0))
    for i in range(n - 2, -1, -1):
        up_to = tiers[i].get("upTo")
        pct = tiers[i].get("pct", 0)
        cmp = "<=" if i == 0 else "<"
        bound = up_to if up_to is not None else 0
        expr = f"IF({ccg_ref}{cmp}{bound},{pct},{expr})"
    return expr


def compute_detail_totals(sheet, salary_table, cql_tiers=None):
    """Tinh Ccg/Cql/Ck/TN/G/VAT/Tong cong theo dung cong thuc va tuy chinh
    lam tron cua tung hang muc du toan chi tiet."""
    days_per_month = salary_table.get("workDaysPerMonth", 26) or 26
    levels = {lvl["id"]: lvl for lvl in salary_table.get("levels", [])}

    def daily_rate(level_id):
        lvl = levels.get(level_id)
        if not lvl:
            return 0
        # Lam tron don gia ngay cong den hang don vi (dong) ngay tai nguon, de moi con so
        # lien quan (Thanh tien tung dong, Ccg, Cql, TN, G, VAT...) deu tinh nhat quan tu
        # CUNG 1 gia tri da lam tron - tranh truong hop nhan tay 3 cot hien thi ra so le
        # so voi cot Thanh tien do dung so chua lam tron o hau truong.
        return round(lvl.get("monthlyRate", 0) / days_per_month)

    if sheet.get("ccgMode") == "lumpsum":
        Ccg = sheet.get("ccgLumpsum", 0) or 0
    else:
        analysis = analyze_expert_rows(sheet.get("expertRows", []), daily_rate)
        Ccg = sum(a["contributes"] for a in analysis)

    if sheet.get("ckMode") == "lumpsum":
        Ck = sheet.get("ckLumpsum", 0) or 0
    else:
        Ck = sum(
            (r.get("quantity") or 0) * (r.get("unitPrice") or 0)
            for r in sheet.get("otherRows", [])
        )

    cql_pct = resolve_cql_pct(Ccg, cql_tiers)
    Cql = Ccg * cql_pct / 100
    TN = (Ccg + Cql + Ck) * (sheet.get("tnPct", 0) or 0) / 100
    G = Ccg + Cql + Ck + TN
    VAT = G * (sheet.get("vatPct", 0) or 0) / 100
    before_discount = G + VAT
    # Linh hoat: giam gia hang muc co the la theo % (mac dinh) hoac so tien co dinh - clamp so tien
    # co dinh de khong vuot qua before_discount (tranh am tong tien). (Dong bo voi computeDetailTotal()
    # trong ui.html.)
    if sheet.get("discountEnabled"):
        if sheet.get("discountMode") == "fixed":
            discount = min(sheet.get("discountFixed", 0) or 0, before_discount)
        else:
            discount = before_discount * (sheet.get("discountPct", 0) or 0) / 100
    else:
        discount = 0
    total = before_discount - discount

    mode = sheet.get("roundingMode", "floor")
    round_enabled = sheet.get("roundEnabled", True)
    # Don vi lam tron: theo lua chon cua nguoi dung khi bat "Ap dung lam tron tong cong";
    # khi TAT, van lam tron ve don vi VND (unit=1) - tien VND khong co don vi nho hon 1 dong,
    # tranh de so le duoi 1 dong phat sinh do nhan/chia phan tram roi cong don ve sau.
    # (Dong bo tuyet doi voi computeDetailTotal() trong ui.html.)
    unit = (sheet.get("roundingUnit", 100000) or 1) if round_enabled else 1

    def _round_to_unit(value):
        if mode == "ceil":
            return math.ceil(value / unit) * unit
        elif mode == "nearest":
            return round(value / unit) * unit
        else:
            return math.floor(value / unit) * unit

    rounded_total = _round_to_unit(total)

    # Danh rieng cho bang bao gia chinh khi bat "Hien thi chi tiet VAT" va bang "Tong hop chi phi":
    # giam gia tru NGAY tren gia tri cong truoc thue (G), roi moi tinh VAT tren phan da giam - ve
    # mat toan hoc cho cung 1 ket qua nhu cach tinh G+VAT roi moi giam gia o tren.
    # Quy giam gia (du la % hay so tien co dinh) ve 1 "ty le giam gia hieu dung" tren before_discount,
    # roi ap dung deu ty le do len G - tong quat hoa cong thuc % goc cho ca truong hop so tien co dinh.
    effective_discount_frac = (discount / before_discount) if (sheet.get("discountEnabled") and before_discount > 0) else 0
    main_vat_base_raw = G * (1 - effective_discount_frac) if sheet.get("discountEnabled") else G
    # Lam tron CA "gia tri truoc thue" VA "gia tri bao gom thue" ve CUNG don vi/kieu lam tron o tren,
    # roi suy ra "gia tri thue" = hieu cua 2 so da lam tron - dam bao tuyet doi
    # truoc thue + thue = bao gom thue, khong con lech do lam tron rieng le tung so
    # (ap dung luon ca khi TAT "Ap dung lam tron tong cong", vi unit khi do = 1 dong).
    main_vat_base = _round_to_unit(main_vat_base_raw)
    main_vat_amount = rounded_total - main_vat_base

    return {
        "Ccg": Ccg, "Ck": Ck, "Cql": Cql, "cqlPct": cql_pct, "TN": TN, "G": G, "VAT": VAT,
        "beforeDiscount": before_discount, "discount": discount, "total": total,
        "roundedTotal": rounded_total, "roundEnabled": round_enabled, "daily_rate": daily_rate,
        "expert_analysis": analyze_expert_rows(sheet.get("expertRows", []), daily_rate),
        "mainVatBase": main_vat_base, "mainVatAmount": main_vat_amount,
    }


def _safe_sheet_name(raw_name, used_names, prefix=""):
    base = "".join(c if (c.isalnum() or c in " _-") else "_" for c in (raw_name or "Hang_muc"))
    name = (prefix + base)[:31]
    final, i = name, 1
    while final in used_names:
        i += 1
        suffix = f"_{i}"
        final = (name[: 31 - len(suffix)] + suffix)
    used_names.add(final)
    return final


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------

def _plan_detail_sheet_excel(detail, totals, show_vat, cql_tiers):
    """Tinh TRUOC (khong ghi o dau vao sheet) toan bo cau truc + CONG THUC cua 1 sheet "Du toan chi tiet"
    trong file Excel: vi tri (so dong) cua tung o quan trong (Ccg/Cql/Ck/TN/CONG/.../TONG CONG LAM TRON)
    va cong thuc thuc su se ghi vao o do - de (1) sheet "Bao gia" chinh co the tham chieu CHEO SANG cac o
    nay bang cong thuc ngay ca khi sheet chi tiet chua duoc ghi, va (2) ham ghi sheet chi tiet (phia duoi)
    chi viec "phat" (deploy) dung cac cong thuc da tinh san nay vao dung o, khong phai tu tinh lai.

    Nguyen tac chung: MOI o so lieu deu la CONG THUC tham chieu nguoc ve cac o dau vao goc (khoi luong,
    ngay cong, don gia...) thay vi ghi san ket qua da tinh - dung y "kem theo cong thuc goc" nguoi dung
    yeu cau, giup ho tu kiem tra/sua truc tiep tren Excel ma khong phai tin tuong mu quang vao 1 con so."""
    tn_pct = detail.get("tnPct", 0) or 0
    vat_pct = detail.get("vatPct", 0) or 0
    round_mode = detail.get("roundingMode", "floor")
    round_unit = (detail.get("roundingUnit", 100000) or 1) if totals["roundEnabled"] else 1
    round_fn = {"ceil": "CEILING", "nearest": "MROUND"}.get(round_mode, "FLOOR")
    discount_enabled = bool(detail.get("discountEnabled"))
    discount_mode = detail.get("discountMode")
    discount_pct = detail.get("discountPct", 0) or 0
    discount_fixed = detail.get("discountFixed", 0) or 0

    ccg_explain = "Trọn gói" if detail.get("ccgMode") == "lumpsum" else "Tổng hợp chi phí chuyên gia"
    ck_explain = "Trọn gói" if detail.get("ckMode") == "lumpsum" else "Tổng hợp chi phí khác"
    round_label = "TỔNG CỘNG (LÀM TRÒN)" if totals["roundEnabled"] else "TỔNG CỘNG"

    rows = [
        {"key": "ccg", "stt": "I", "label": "CHI PHÍ CHUYÊN GIA", "sym": "Ccg", "explain": ccg_explain, "bold": False},
        {"key": "cql", "stt": "II", "label": "CHI PHÍ QUẢN LÝ", "sym": "Cql", "explain": f"({totals['cqlPct']}%)*Ccg", "bold": False},
        {"key": "ck", "stt": "III", "label": "CHI PHÍ KHÁC", "sym": "Ck", "explain": ck_explain, "bold": False},
        {"key": "tn", "stt": "IV", "label": "THU NHẬP CHỊU THUẾ TÍNH TRƯỚC", "sym": "TN", "explain": f"{tn_pct}%*(Ccg+Cql+Ck)", "bold": False},
        {"key": "cong", "stt": "", "label": "CỘNG", "sym": "G", "explain": "Ccg+Cql+Ck+TN", "bold": True},
    ]
    if show_vat:
        discount_desc = "" if discount_mode == "fixed" else f"{discount_pct}%"
        round_before_tax_label = "GIÁ TRỊ TRƯỚC THUẾ (LÀM TRÒN)" if totals["roundEnabled"] else "GIÁ TRỊ TRƯỚC THUẾ"
        if discount_enabled:
            round_before_tax_label_disc = "GIÁ TRỊ TRƯỚC THUẾ (SAU GIẢM GIÁ) (LÀM TRÒN)" if totals["roundEnabled"] else "GIÁ TRỊ TRƯỚC THUẾ (SAU GIẢM GIÁ)"
            rows.append({"key": "discount_display", "stt": "", "label": "Giảm giá dịch vụ", "sym": "", "explain": discount_desc, "bold": False})
            rows.append({"key": "vat_base", "stt": "", "label": round_before_tax_label_disc, "sym": "G'", "explain": "G - Giảm giá dịch vụ", "bold": True})
            vat_explain = f"{vat_pct}% x G'"
        else:
            rows.append({"key": "vat_base", "stt": "", "label": round_before_tax_label, "sym": "G", "explain": "(làm tròn)", "bold": True})
            vat_explain = f"{vat_pct}% x G"
        rows.append({"key": "vat", "stt": "V", "label": "Thuế giá trị gia tăng", "sym": "GTGT", "explain": vat_explain, "bold": False})
    else:
        rows.append({"key": "vat", "stt": "V", "label": "Thuế giá trị gia tăng", "sym": "GTGT", "explain": f"{vat_pct}% x G", "bold": False})
        rows.append({"key": "before_discount", "stt": "", "label": "TỔNG CỘNG", "sym": "", "explain": "G + GTGT", "bold": True})
        if discount_enabled:
            discount_desc = "" if discount_mode == "fixed" else f"{discount_pct}%"
            rows.append({"key": "discount_display", "stt": "", "label": "Giảm giá dịch vụ", "sym": "", "explain": discount_desc, "bold": False})
    rows.append({"key": "rounded_total", "stt": "", "label": round_label, "sym": "", "explain": "", "bold": True})

    header_row = 3
    first_data_row = header_row + 1
    row_of = {rd["key"]: first_data_row + i for i, rd in enumerate(rows)}
    summary_last_row = first_data_row + len(rows) - 1

    def E(key):
        return f"E{row_of[key]}"

    formulas = {}
    formulas["cql"] = f"={E('ccg')}*({_cql_pct_formula(E('ccg'), cql_tiers)})/100"
    formulas["tn"] = f"=({E('ccg')}+{E('cql')}+{E('ck')})*{tn_pct}/100"
    formulas["cong"] = f"={E('ccg')}+{E('cql')}+{E('ck')}+{E('tn')}"

    if show_vat:
        if discount_enabled:
            if discount_mode == "fixed":
                before_disc_expr = f"{E('cong')}*(1+{vat_pct}/100)"
                vat_base_raw = f"{E('cong')}*(1-MIN({discount_fixed},{before_disc_expr})/({before_disc_expr}))"
                total_raw = f"{before_disc_expr}-MIN({discount_fixed},{before_disc_expr})"
            else:
                vat_base_raw = f"{E('cong')}*(1-{discount_pct}/100)"
                total_raw = f"{E('cong')}*(1+{vat_pct}/100)*(1-{discount_pct}/100)"
            formulas["vat_base"] = f"={round_fn}({vat_base_raw},{round_unit})"
            formulas["discount_display"] = f"=-({E('cong')}-{E('vat_base')})"
        else:
            formulas["vat_base"] = f"={round_fn}({E('cong')},{round_unit})"
            total_raw = f"{E('cong')}*(1+{vat_pct}/100)"
        formulas["rounded_total"] = f"={round_fn}({total_raw},{round_unit})"
        formulas["vat"] = f"={E('rounded_total')}-{E('vat_base')}"
    else:
        formulas["vat"] = f"={E('cong')}*{vat_pct}/100"
        formulas["before_discount"] = f"={E('cong')}+{E('vat')}"
        if discount_enabled:
            if discount_mode == "fixed":
                formulas["discount_display"] = f"=-MIN({discount_fixed},{E('before_discount')})"
            else:
                formulas["discount_display"] = f"=-{E('before_discount')}*{discount_pct}/100"
            formulas["rounded_total"] = f"={round_fn}({E('before_discount')}+{E('discount_display')},{round_unit})"
        else:
            formulas["rounded_total"] = f"={round_fn}({E('before_discount')},{round_unit})"

    # --- Len ke hoach vi tri 2 bang duoi (Dien giai chi phi chuyen gia / chi phi khac) - dua theo DUNG
    # trinh tu cac dong `r += 1`/lap header/lap dong "Tong hop..." nhu ham ghi thuc te ben duoi, de vi
    # tri tinh truoc va vi tri ghi thuc te LUON khop nhau tuyet doi.
    pointer = summary_last_row + 2  # qua 1 dong trong

    expert_layout = None
    if detail.get("ccgMode") != "lumpsum":
        expert_rows_list = detail.get("expertRows", [])
        analysis = totals["expert_analysis"]
        title_row = pointer
        header_row_e = pointer + 1
        summary_row_e = pointer + 2
        first_data_row_e = pointer + 3
        # row_of_e: vi tri (so dong Excel) cua TUNG dong expertRows (du la dong "me" co dong con hay
        # dong "la" - 1 muc chuyen gia); children_rows_e: voi moi dong "me", danh sach cac so dong Excel
        # cua CAC DONG CON TRUC TIEP (khong phai chau) - de phia duoi dung xay cong thuc
        # SUM(G{r1},G{r2},...) THAM CHIEU DICH DANH tung dong con, dung tai moi cap do sau bao nhieu,
        # thay vi 1 pham vi lien tiep (chi dung voi dung 2 cap).
        row_of_e = {}
        row_counts = []
        r_ptr = first_data_row_e
        for i, er in enumerate(expert_rows_list):
            row_of_e[i] = r_ptr
            if analysis[i]["hasChildren"]:
                row_counts.append(1)
                r_ptr += 1
            else:
                lv = er.get("levels") or [{"levelId": er.get("levelId"), "quantity": er.get("quantity"), "days": er.get("days")}]
                cnt = max(1, len(lv))
                row_counts.append(cnt)
                r_ptr += cnt
        children_rows_e = {
            i: [row_of_e[ci] for ci in analysis[i]["childIndices"]]
            for i in range(len(expert_rows_list))
        }
        total_rows = sum(row_counts)
        last_row_e = first_data_row_e + total_rows - 1 if total_rows else first_data_row_e - 1
        expert_layout = {
            "title_row": title_row, "header_row": header_row_e, "summary_row": summary_row_e,
            "first_data_row": first_data_row_e, "last_row": last_row_e,
            "row_of": row_of_e, "children_rows": children_rows_e,
        }
        pointer = last_row_e + 1
    pointer += 1  # dong trong thu 2 - LUON chay (giong het ban goc, khong phu thuoc ccgMode)

    other_layout = None
    if detail.get("ckMode") != "lumpsum":
        other_rows_list = detail.get("otherRows", [])
        n_other = len(other_rows_list)
        title_row = pointer
        header_row_o = pointer + 1
        summary_row_o = pointer + 2
        first_data_row_o = pointer + 3
        last_row_o = first_data_row_o + n_other - 1 if n_other else first_data_row_o - 1
        other_layout = {
            "title_row": title_row, "header_row": header_row_o, "summary_row": summary_row_o,
            "first_data_row": first_data_row_o, "last_row": last_row_o,
        }

    # --- Dien cong thuc Ccg (SUM cac dong "Thanh tien" cua bang Dien giai chi phi chuyen gia, TRU cac
    # dong "me" co dong con - vi dong me chi hien THAM KHAO tong cac dong con, khong duoc cong thêm lan
    # nua, giong het quy tac analyze_expert_rows()/analyzeExpertRows() o tren/trong ui.html) ----
    if detail.get("ccgMode") == "lumpsum":
        formulas["ccg"] = totals["Ccg"]
    elif expert_layout and expert_layout["last_row"] >= expert_layout["first_data_row"]:
        expert_rows_list = detail.get("expertRows", [])
        analysis = totals["expert_analysis"]
        # Cong thuc nay dung o BAT KY so cap long nhau nao, khong can sua lai du co 2, 3 hay nhieu cap:
        # SUM(toan bo pham vi) - SUM(cac dong "me") = SUM(cac dong "la") - vi gia tri cua MOI dong me
        # (duoc ghi vao cot G phia duoi) luon dung bang tong cac dong con TRUC TIEP cua no (de quy), nen
        # tru het cac dong me di se chi con lai dung phan cua cac dong la - khop 100% voi quy tac
        # "contributes" trong analyze_expert_rows()/analyzeExpertRows().
        parent_rows = [expert_layout["row_of"][i] for i in range(len(expert_rows_list)) if analysis[i]["hasChildren"]]
        ccg_formula = f"=SUM(G{expert_layout['first_data_row']}:G{expert_layout['last_row']})"
        for pr in parent_rows:
            ccg_formula += f"-G{pr}"
        formulas["ccg"] = ccg_formula
    else:
        formulas["ccg"] = 0

    # --- Dien cong thuc Ck (SUM toan bo cot "Thanh tien" cua bang Dien giai chi phi khac - KHONG loai
    # tru dong con nao, vi Ck von di da cong TAT CA otherRows khong phan biet me/con - xem compute_detail_totals) ----
    if detail.get("ckMode") == "lumpsum":
        formulas["ck"] = totals["Ck"]
    elif other_layout and other_layout["last_row"] >= other_layout["first_data_row"]:
        formulas["ck"] = f"=SUM(F{other_layout['first_data_row']}:F{other_layout['last_row']})"
    else:
        formulas["ck"] = 0

    return {
        "rows": rows, "row_of": row_of, "header_row": header_row, "first_data_row": first_data_row,
        "summary_last_row": summary_last_row, "formulas": formulas,
        "expert_layout": expert_layout, "other_layout": other_layout,
    }


def export_excel_file(state, save_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bao gia"

    bold = Font(bold=True, name="Times New Roman", size=12)
    normal = Font(name="Times New Roman", size=12)
    italic_normal = Font(italic=True, name="Times New Roman", size=12)
    title_font = Font(bold=True, name="Times New Roman", size=20)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    show_vat = bool(state.get("showVatBreakdown"))
    last_col = 5 if show_vat else 3  # cot cuoi cung cua bang bao gia chinh - dung de canh cac khoi khac
    last_col_letter = get_column_letter(last_col)

    show_header_logo = state.get("showHeaderLogo", True)
    logo_bytes = _decode_logo_bytes(state.get("logoDataUri")) if show_header_logo else None
    footer_caption = _resolve_footer_caption(state)

    salary_table = state.get("salaryTable", {"levels": [], "workDaysPerMonth": 26})
    levels_by_id = {lvl["id"]: lvl for lvl in salary_table.get("levels", [])}
    cql_tiers = state.get("cqlTiers")
    used_names = {ws.title}

    # Ten sheet "Bang luong" + vi tri (dong) cua tung Muc trong do - DANH TRUOC (chua ghi gi ca), de cot
    # "Don gia ngay cong" trong bang Dien giai chi phi chuyen gia cua moi hang muc co the THAM CHIEU CHEO
    # SANG DUNG dong cua sheet nay bang cong thuc, thay vi ghi san 1 con so da tinh.
    salary_header_row = 5
    level_row_of = {lvl["id"]: salary_header_row + 1 + i for i, lvl in enumerate(salary_table.get("levels", []))}
    any_sheet_uses_salary_table = any(d.get("ccgMode") == "linked" for d in state.get("detailSheets", []))
    sal_sheet_name = _safe_sheet_name("Bang luong", used_names) if any_sheet_uses_salary_table else None
    days_per_month = salary_table.get("workDaysPerMonth", 26) or 26

    # ---- Buoc 1: len ke hoach (KHONG ghi gi vao workbook) cho tung hang muc du toan chi tiet - biet
    # truoc ten sheet Excel + toan bo cong thuc + vi tri cac o quan trong, de sheet "Bao gia" chinh (ghi
    # NGAY SAU day) co the tham chieu thang bang cong thuc sang cac sheet nay du chung chua duoc tao. ----
    detail_plans = []
    for detail in state.get("detailSheets", []):
        totals = compute_detail_totals(detail, salary_table, cql_tiers)
        raw_name = detail.get("name") or "Hang muc"
        sheet_name = _safe_sheet_name(raw_name, used_names, prefix="DT_")
        plan = _plan_detail_sheet_excel(detail, totals, show_vat, cql_tiers)
        plan.update({"id": detail.get("id"), "sheet_name": sheet_name, "raw_name": raw_name, "totals": totals})
        detail_plans.append(plan)
    plan_by_id = {p["id"]: p for p in detail_plans}

    def _amount_ref(plan, key):
        return f"='{plan['sheet_name']}'!E{plan['row_of'][key]}"

    row = 1
    logo_inserted = False
    if logo_bytes:
        try:
            xl_img = XLImage(BytesIO(logo_bytes))
            target_h = 44
            if xl_img.height:
                scale = target_h / xl_img.height
                xl_img.width = max(1, round(xl_img.width * scale))
            xl_img.height = target_h
            ws.add_image(xl_img, f"A{row}")
            ws.row_dimensions[row].height = 34
            logo_inserted = True
        except Exception:
            logo_inserted = False
    if not logo_inserted and show_header_logo:
        # Bat "hien logo" nhung khong co/khong doc duoc anh -> fallback ve chu ITD Vietnam mac dinh.
        # Neu nguoi dung tat han "hien logo" thi khong hien gi ca (dung y "an logo" that su).
        ws.cell(row=row, column=1, value="ITD Vietnam").font = Font(bold=True, size=14, name="Times New Roman")
    ws.cell(row=row, column=3, value="BAO GIA").font = title_font
    row += 1
    ws.cell(row=row, column=1, value=f"Ma so: {state.get('quoteCode','')}").font = bold
    row += 2

    ws.cell(row=row, column=1, value=state.get("companyName", "")).font = bold
    ws.cell(row=row, column=last_col, value=f"Tel: {state.get('companyPhone','')}").font = normal
    row += 1
    ws.cell(row=row, column=1, value=state.get("companyAddr", "")).font = normal
    ws.cell(row=row, column=last_col, value=f"Email: {state.get('companyEmail','')}").font = normal
    row += 2

    ws.cell(row=row, column=1, value=f"Kinh gui: {state.get('recipient','')}").font = bold
    row += 1
    intro_text = f"{state.get('introPrefix','')} {state.get('fixedClause','')}".strip()
    ws.cell(row=row, column=1, value=intro_text).font = normal
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.row_dimensions[row].height = 40
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    row += 2

    ws.cell(row=row, column=last_col, value="Don vi tinh: VND").font = Font(italic=True, size=10, name="Times New Roman")
    ws.cell(row=row, column=last_col).alignment = right
    row += 1

    header_row = row
    if show_vat:
        headers = ["TT", "Noi dung", "Don gia truoc VAT", "VAT", "Don gia sau VAT"]
    else:
        headers = ["TT", "Noi dung", "Gia dich vu"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = bold
        c.alignment = center
        c.border = border
    row += 1

    items = state.get("items", [])
    item_amount_rows = []  # dong cua tung o "Gia dich vu"/"Don gia sau VAT" (cot last_col) - dung SUM
    # lai chinh cac o nay cho dong "Tong cong gia tri dich vu" phia duoi, thay vi 1 con so tinh san.
    for idx, it in enumerate(items):
        is_main = bool(it.get("isMain"))
        # Dong hang muc chinh co the duoc nguoi dung bat tuy chon "cong tong cac hang muc con vao
        # dong nay" (xem g.sumSubItems/isSumRow ben ui.html) - khi do no CUNG hien gia tri nhu 1 dong
        # binh thuong (khong merge/de trong nua), va gia tri la CONG THUC SUM tham chieu thang toi cac
        # dong hang muc con ngay phia duoi no (vi tri cac dong nay xac dinh truoc duoc vi cau truc
        # items luon xep hang muc con lien tiep ngay sau hang muc chinh cua no).
        is_sum_row = is_main and bool(it.get("isSumRow"))
        # Dong hang muc chinh: in dam toan dong. Dong hang muc con: chu thang (mac dinh) hoac in nghieng
        # neu nguoi dung bat "In nghieng cac dong hang muc con" (xem g.subItemsItalic/row.italic ben ui.html).
        f = bold if is_main else (italic_normal if it.get("italic") else normal)
        ws.cell(row=row, column=1, value=it.get("number", "")).font = f
        ws.cell(row=row, column=1).alignment = center
        ws.cell(row=row, column=2, value=it.get("name", "")).font = f

        if is_main and not is_sum_row:
            if show_vat:
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
            # show_vat False: khong lam gi them, cot "Gia dich vu" de trong - giong het ui.html.
        else:
            n_children = 0
            if is_sum_row:
                # Dung DUNG so luong hang muc con ma ui.html da tinh san va gui kem (childCount) -
                # KHONG tu do quet cac dong lien tiep co isMain=false phia sau, vi ngay sau nhom nay
                # co the la 1 hang muc rieng le/nhom showMain=false khac, cung co isMain=false nhung
                # KHONG phai con cua no (quet kieu do se SUM nham sang ca hang muc khong lien quan).
                n_children = int(it.get("childCount") or 0)
                child_start, child_end = row + 1, row + n_children
                plan = None
            else:
                plan = plan_by_id.get(it.get("detailId"))

            if show_vat:
                bv = it.get("beforeVat", None)
                va = it.get("vatAmount", None)
                if bv is not None:
                    if is_sum_row:
                        bv_val = f"=SUM(C{child_start}:C{child_end})" if n_children else round(bv)
                    else:
                        bv_val = _amount_ref(plan, "vat_base") if (plan and "vat_base" in plan["row_of"]) else round(bv)
                    c3 = ws.cell(row=row, column=3, value=bv_val)
                    c3.number_format = "#,##0"
                    c3.font = f
                if va is not None:
                    if is_sum_row:
                        va_val = f"=SUM(D{child_start}:D{child_end})" if n_children else round(va)
                    else:
                        va_val = _amount_ref(plan, "vat") if (plan and "vat" in plan["row_of"]) else round(va)
                    c4 = ws.cell(row=row, column=4, value=va_val)
                    c4.number_format = "#,##0"
                    c4.font = f
            amt = it.get("amount", None)
            if amt is not None:
                if is_sum_row:
                    # Don gia hang muc chinh (khi bat "cong tong") = CONG THUC SUM cac dong hang muc
                    # con ngay phia duoi, KHONG cong lai vao item_amount_rows (tranh cong trung 2 lan
                    # vao dong "Tong cong gia tri dich vu" - cac hang muc con da tu cong roi).
                    amt_val = f"=SUM({last_col_letter}{child_start}:{last_col_letter}{child_end})" if n_children else round(amt)
                else:
                    # Don gia hang muc nay LUON tham chieu (bang cong thuc) toi dung o "TONG CONG (LAM TRON)"
                    # cua chinh sheet du toan chi tiet hang muc do - dam bao nguoi dung sua so lieu trong sheet
                    # chi tiet la bang bao gia chinh tu cap nhat theo, khong con phai go lai tay o day.
                    amt_val = _amount_ref(plan, "rounded_total") if plan else amt
                c5 = ws.cell(row=row, column=last_col, value=amt_val)
                c5.number_format = "#,##0"
                c5.font = f
                if not is_sum_row:
                    item_amount_rows.append(row)
        for col in range(1, last_col + 1):
            ws.cell(row=row, column=col).border = border
        row += 1

    mode = state.get("mode", "single")
    if mode == "single":
        subtotal_expr = "+".join(f"{last_col_letter}{r}" for r in item_amount_rows) or "0"
        discount_row = None
        if state.get("discountEnabled"):
            pct = state.get("discountPct", 0)
            discount_label = "Giam gia dich vu" if state.get("discountMode") == "fixed" else f"Giam gia dich vu {pct}%"
            if state.get("discountMode") == "fixed":
                fixed_amt = state.get("discountFixed", 0) or 0
                discount_formula = f"=-MIN({fixed_amt},({subtotal_expr}))"
            else:
                discount_formula = f"=-({subtotal_expr})*{pct}/100"
            ws.cell(row=row, column=1).border = border
            c2 = ws.cell(row=row, column=2, value=discount_label)
            c2.font = Font(italic=True, name="Times New Roman", size=12)
            c2.border = border
            if show_vat:
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col - 1)
            c_amt = ws.cell(row=row, column=last_col, value=discount_formula)
            c_amt.number_format = "#,##0"
            c_amt.border = border
            for col in range(1, last_col + 1):
                ws.cell(row=row, column=col).border = border
            discount_row = row
            row += 1
        total_formula = f"=({subtotal_expr})+{last_col_letter}{discount_row}" if discount_row else f"={subtotal_expr}"
        ws.cell(row=row, column=1).border = border
        c2 = ws.cell(row=row, column=2, value="Tong cong gia tri dich vu")
        c2.font = bold
        c2.border = border
        if show_vat:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col - 1)
        c_amt = ws.cell(row=row, column=last_col, value=total_formula)
        c_amt.font = bold
        c_amt.number_format = "#,##0"
        for col in range(1, last_col + 1):
            ws.cell(row=row, column=col).border = border
        row += 1

    row += 1
    for note in state.get("footerNotes", []):
        ws.cell(row=row, column=1, value=note).font = normal
        row += 1

    row += 1
    closing = ws.cell(row=row, column=1, value=state.get("closingNote", ""))
    closing.font = bold
    closing.alignment = center
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    row += 2

    date_line = f"Ha Noi, ngay {state.get('qDay','')} thang {state.get('qMonth','')} nam {state.get('qYear','')}"
    ws.cell(row=row, column=last_col, value=date_line).alignment = right
    ws.cell(row=row, column=last_col).font = normal
    row += 1
    sign_company_name = (state.get("signCompanyName") or state.get("companyName", "") or "").upper()
    ws.cell(row=row, column=last_col, value=sign_company_name).font = bold
    ws.cell(row=row, column=last_col).alignment = right

    ws.column_dimensions["A"].width = 8
    if show_vat:
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 14
        ws.column_dimensions["E"].width = 18
    else:
        ws.column_dimensions["B"].width = 55
        ws.column_dimensions["C"].width = 22

    # Chan trang khi in (chi hien trong Print Preview / ban in, khong hien o che do xem thuong)
    ws.oddFooter.left.text = footer_caption
    ws.oddFooter.left.size = 9
    ws.oddFooter.left.font = "Times New Roman,Bold"

    # ---- Buoc 2: ghi thuc su tung sheet "Du toan chi tiet" - PHAT (deploy) dung cac cong thuc + vi tri
    # da tinh san o plan (buoc 1) vao dung o, khong tu tinh lai gi them. ----
    for detail, plan in zip(state.get("detailSheets", []), detail_plans):
        totals = plan["totals"]
        daily_rate = totals["daily_rate"]
        raw_name = plan["raw_name"]
        sheet_name = plan["sheet_name"]
        row_of = plan["row_of"]
        formulas = plan["formulas"]

        dws = wb.create_sheet(sheet_name)
        dws.cell(row=1, column=1, value=raw_name).font = Font(bold=True, size=14, name="Times New Roman")

        # Bang Tong hop chi phi
        for col, h in enumerate(["STT", "Khoan muc chi phi", "Ky hieu", "Dien giai", "Thanh tien"], start=1):
            c = dws.cell(row=plan["header_row"], column=col, value=h)
            c.font = bold
            c.border = border
            c.alignment = center

        for i, rd in enumerate(plan["rows"]):
            r = plan["first_data_row"] + i
            f = bold if rd["bold"] else normal
            dws.cell(row=r, column=1, value=rd["stt"]).font = f
            dws.cell(row=r, column=1).alignment = center
            dws.cell(row=r, column=2, value=rd["label"]).font = f
            dws.cell(row=r, column=3, value=rd["sym"]).font = f
            dws.cell(row=r, column=3).alignment = center
            dws.cell(row=r, column=4, value=rd["explain"]).font = f
            value = formulas.get(rd["key"])
            c5 = dws.cell(row=r, column=5, value=value if isinstance(value, str) else round(value))
            c5.number_format = "#,##0"
            c5.font = f
            for col in range(1, 6):
                dws.cell(row=r, column=col).border = border

        # Dien giai chi phi chuyen gia
        el = plan["expert_layout"]
        if el:
            dws.cell(row=el["title_row"], column=1, value="Dien giai chi phi chuyen gia").font = bold
            for col, h in enumerate(
                ["STT", "Cong viec thuc hien", "Muc chuyen gia", "Khoi luong chuyen gia",
                 "Ngay cong", "Don gia ngay cong", "Thanh tien (VND)"], start=1
            ):
                c = dws.cell(row=el["header_row"], column=col, value=h)
                c.font = bold
                c.border = border
                c.alignment = center
            r = el["summary_row"]
            dws.cell(row=r, column=2, value="Tong hop chi phi chuyen gia tu van").font = bold
            # O nay chi la "bi danh" (alias) cua chinh o Ccg trong bang Tong hop chi phi o tren - luon
            # tu dong khop nhau, khong bao gio lech du sua o dau.
            c7 = dws.cell(row=r, column=7, value=f"=E{row_of['ccg']}")
            c7.number_format = "#,##0"
            c7.font = bold
            for col in range(1, 8):
                dws.cell(row=r, column=col).border = border

            expert_rows_list = detail.get("expertRows", [])
            nums = compute_expert_row_numbers(expert_rows_list)
            analysis = totals["expert_analysis"]
            r = el["first_data_row"]
            for i, er in enumerate(expert_rows_list):
                a = analysis[i]
                depth = a.get("depth", 0) or 0
                desc_wrap = Alignment(wrap_text=True, vertical="center", indent=depth)
                if a["hasChildren"]:
                    # Dong "me" (co dong con) CHI hien TONG cac dong con TRUC TIEP cua no (khong phai
                    # tat ca chau) - dung cong thuc SUM(G{r1},G{r2},...) tham chieu DICH DANH tung dong
                    # con lay tu children_rows (da tinh san o _plan_detail_sheet_excel, dua tren cau
                    # truc depth thuc te) thay vi 1 pham vi lien tiep (chi dung o dung 2 cap).
                    child_rows = el["children_rows"].get(i, [])
                    dws.cell(row=r, column=1, value=nums[i]).alignment = center
                    c2 = dws.cell(row=r, column=2, value=er.get("desc", ""))
                    c2.alignment = desc_wrap
                    dws.cell(row=r, column=3, value="-").alignment = center
                    dws.cell(row=r, column=4, value="-").alignment = center
                    dws.cell(row=r, column=5, value="-").alignment = center
                    dws.cell(row=r, column=6, value="-").alignment = center
                    if child_rows:
                        refs = ",".join(f"G{cr}" for cr in child_rows)
                        c7 = dws.cell(row=r, column=7, value=f"=SUM({refs})")
                    else:
                        c7 = dws.cell(row=r, column=7, value=0)
                    c7.number_format = "#,##0"
                    c7.font = bold
                    for col in range(1, 8):
                        dws.cell(row=r, column=col).border = border
                    r += 1
                else:
                    # Moi muc chuyen gia trong cung 1 cong viec xuong 1 dong rieng; cot STT va
                    # Cong viec thuc hien duoc gop (merge) doc theo nhom cac dong nay. Thanh tien =
                    # Khoi luong x Ngay cong x Don gia ngay cong (cong thuc that, khong ghi san so),
                    # Don gia ngay cong lai THAM CHIEU CHEO sang chinh sheet Bang luong chuyen gia.
                    row_levels = er.get("levels") or [{
                        "levelId": er.get("levelId"), "quantity": er.get("quantity"), "days": er.get("days"),
                    }]
                    group_start = r
                    for lv in row_levels:
                        lvl = levels_by_id.get(lv.get("levelId"))
                        qty = lv.get("quantity") or 0
                        days = lv.get("days") or 0
                        dws.cell(row=r, column=3, value=lvl.get("name") if lvl else "").alignment = center
                        c4 = dws.cell(row=r, column=4, value=qty)
                        c4.alignment = center
                        c5 = dws.cell(row=r, column=5, value=days)
                        c5.alignment = center
                        lvl_row = level_row_of.get(lv.get("levelId"))
                        if sal_sheet_name and lvl_row:
                            c6 = dws.cell(row=r, column=6, value=f"='{sal_sheet_name}'!C{lvl_row}")
                        else:
                            c6 = dws.cell(row=r, column=6, value=round(daily_rate(lv.get("levelId"))))
                        c6.number_format = "#,##0"
                        c7 = dws.cell(row=r, column=7, value=f"=D{r}*E{r}*F{r}")
                        c7.number_format = "#,##0"
                        for col in range(1, 8):
                            dws.cell(row=r, column=col).border = border
                        r += 1
                    group_end = r - 1
                    dws.cell(row=group_start, column=1, value=nums[i]).alignment = center
                    c2 = dws.cell(row=group_start, column=2, value=er.get("desc", ""))
                    c2.alignment = desc_wrap
                    if group_end > group_start:
                        dws.merge_cells(start_row=group_start, start_column=1, end_row=group_end, end_column=1)
                        dws.merge_cells(start_row=group_start, start_column=2, end_row=group_end, end_column=2)

        # Dien giai chi phi khac
        ol = plan["other_layout"]
        if ol:
            dws.cell(row=ol["title_row"], column=1, value="Dien giai chi phi khac").font = bold
            for col, h in enumerate(["STT", "Vat tu, vat lieu", "Don vi", "Khoi luong", "Don gia (d)", "Thanh tien (d)"], start=1):
                c = dws.cell(row=ol["header_row"], column=col, value=h)
                c.font = bold
                c.border = border
                c.alignment = center
            r = ol["summary_row"]
            dws.cell(row=r, column=2, value="Tong cong").font = bold
            # O nay chi la "bi danh" cua chinh o Ck trong bang Tong hop chi phi o tren.
            c6 = dws.cell(row=r, column=6, value=f"=E{row_of['ck']}")
            c6.number_format = "#,##0"
            c6.font = bold
            for col in range(1, 7):
                dws.cell(row=r, column=col).border = border

            other_rows_list = detail.get("otherRows", [])
            other_nums = compute_row_numbers(other_rows_list)
            r = ol["first_data_row"]
            for i, orow in enumerate(other_rows_list):
                qty = orow.get("quantity") or 0
                price = orow.get("unitPrice") or 0
                dws.cell(row=r, column=1, value=other_nums[i])
                dws.cell(row=r, column=2, value=orow.get("desc", ""))
                dws.cell(row=r, column=3, value=orow.get("unit", ""))
                c4 = dws.cell(row=r, column=4, value=qty)
                c5 = dws.cell(row=r, column=5, value=price)
                c5.number_format = "#,##0"
                # Thanh tien = Khoi luong x Don gia (cong thuc that, tham chieu chinh 2 o vua ghi).
                c6 = dws.cell(row=r, column=6, value=f"=D{r}*E{r}")
                c6.number_format = "#,##0"
                for col in range(1, 7):
                    dws.cell(row=r, column=col).border = border
                r += 1

        dws.column_dimensions["A"].width = 6
        dws.column_dimensions["B"].width = 38
        dws.column_dimensions["C"].width = 14
        dws.column_dimensions["D"].width = 14
        dws.column_dimensions["E"].width = 12
        dws.column_dimensions["F"].width = 16
        dws.column_dimensions["G"].width = 16

        dws.oddFooter.left.text = footer_caption
        dws.oddFooter.left.size = 9
        dws.oddFooter.left.font = "Times New Roman,Bold"

    # ---- Sheet Bang luong chuyen gia (dung chung) ----
    # CHI tao sheet nay khi co IT NHAT 1 hang muc dang dung ccgMode "linked" (Link tu dien giai) - tuc
    # co thuc su tham chieu den bang luong. Neu TAT CA hang muc deu chon "Tron goi" (lumpsum) thi bang
    # luong khong duoc dung den o dau ca, xuat ra se thua. Vi tri (level_row_of) da danh truoc o dau ham
    # de khop tuyet doi voi cong thuc cac sheet chi tiet da tham chieu toi.
    if sal_sheet_name:
        sws = wb.create_sheet(sal_sheet_name)
        sws.cell(row=1, column=1, value="BANG LUONG CHUYEN GIA").font = Font(bold=True, size=14, name="Times New Roman")
        sws.cell(row=2, column=1, value=f"Ap dung theo: {salary_table.get('circularName','')}").font = normal
        sws.cell(row=3, column=1, value=f"So ngay cong/thang: {days_per_month}").font = normal
        for col, h in enumerate(["Muc", "Dinh muc/thang", "Dinh muc/ngay"], start=1):
            c = sws.cell(row=salary_header_row, column=col, value=h)
            c.font = bold
            c.border = border
            c.alignment = center
        for lvl in salary_table.get("levels", []):
            r = level_row_of[lvl["id"]]
            sws.cell(row=r, column=1, value=lvl.get("name", "")).border = border
            c2 = sws.cell(row=r, column=2, value=lvl.get("monthlyRate", 0))
            c2.number_format = "#,##0"
            c2.border = border
            # Dinh muc/ngay = Dinh muc/thang / so ngay cong quy doi - cong thuc that, khong ghi san so;
            # day chinh la o duoc cac sheet du toan chi tiet tham chieu cheo sang cho cot "Don gia ngay cong".
            c3 = sws.cell(row=r, column=3, value=f"=ROUND(B{r}/{days_per_month},0)")
            c3.number_format = "#,##0"
            c3.border = border
        sws.column_dimensions["A"].width = 20
        sws.column_dimensions["B"].width = 18
        sws.column_dimensions["C"].width = 18
        sws.oddFooter.left.text = footer_caption
        sws.oddFooter.left.size = 9
        sws.oddFooter.left.font = "Times New Roman,Bold"

    wb.save(save_path)


# ---------------------------------------------------------------------------
# WORD
# ---------------------------------------------------------------------------

# ---- Cac mau sac/kich thuoc dung chung - DONG BO voi CSS cua ui.html (xem cac bien --accent,
# --accent-soft, mau dai chan trang #d9d4c4, mau hang "CONG"/"TONG CONG" #f0ede4, mau hang tong
# bang bao gia chinh #f5f0e8...) de file Word xuat ra nhin giong het ban PDF/in, khong phai
# chinh lai mau/khung/canh le bang tay.
ACCENT_RGB = RGBColor(0x8A, 0x3B, 0x2D)
ACCENT_SOFT_HEX = "F1E3D8"
FOOTER_BAND_HEX = "D9D4C4"
SUM_ROW_HEX = "F0EDE4"
MAIN_TOTAL_ROW_HEX = "F5F0E8"
FOOTER_CAPTION_RGB = RGBColor(0x1F, 0x3A, 0x5C)


def _set_a4_page(section):
    """Kho giay A4 + le trang dung KHOP voi padding cua .sheet khi in/xuat PDF (20mm/20mm/30mm/25mm =
    top/right/bottom/left) - mac dinh cua python-docx la kho Letter (21.59x27.94cm, le ~2.5-3.2cm), day
    chinh la nguyen nhan file Word xuat ra truoc day bi sai kho giay/le trang, buoc nguoi dung phai tu
    chinh lai Page Setup truoc khi in."""
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(3.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.0)


def _usable_width_cm(section):
    return (section.page_width - section.left_margin - section.right_margin) / 360000


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _shade_row(row, hex_color):
    for cell in row.cells:
        _shade_cell(cell, hex_color)


def _repeat_header_row(row):
    """Lap lai dong tieu de bang o MOI trang khi bang bi tran qua nhieu trang (tblHeader) - giong
    het cach cac bang trang chi tiet trong ban PDF luon co the hien lai dong tieu de."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _prevent_row_split(row):
    """Khong cho 1 dong bang bi cat lam doi giua 2 trang - neu khong du cho, ca dong se
    tu dong chuyen nguyen sang trang sau (tranh hien tuong dong bi ngat ngang xau xi)."""
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _set_table_col_widths_cm(table, widths_cm):
    """Dat do rong CO DINH cho tung cot (khong de Word tu co gian theo noi dung) - dam bao ty le
    cac cot giong voi ban PDF/in thay vi bi bop meo mac dinh cua Word."""
    table.autofit = False
    table.allow_autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_cm):
                cell.width = Cm(widths_cm[idx])
    for idx, col in enumerate(table.columns):
        if idx < len(widths_cm):
            col.width = Cm(widths_cm[idx])


def _set_table_font_size(table, pt):
    """Thu nho font cac bang du lieu (STT/dien giai/thanh tien...) nho hon than chu chinh van (Pt(13))
    giong ty le trong CSS (table.quote-table font-size ~12.5px, extra-page table ~12px so voi body 13px)
    - vua giong ban PDF hon, vua tranh cac tieu de cot ngan nhu "STT" bi be chu xuong dong xau xi."""
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(pt)


def _style_data_table(table, col_widths_cm, header_row_idx=0, header_fill=ACCENT_SOFT_HEX):
    """Ap dung dong bo cho MOI bang du lieu (bang bao gia chinh, Tong hop chi phi, Dien giai chi
    phi chuyen gia/khac, Bang luong): do rong cot co dinh, to mau dong tieu de, lap lai dong tieu
    de o trang sau neu bang bi tran trang, va khong cho 1 dong bi cat doi giua 2 trang."""
    _set_table_col_widths_cm(table, col_widths_cm)
    _shade_row(table.rows[header_row_idx], header_fill)
    _repeat_header_row(table.rows[header_row_idx])
    for row in table.rows:
        _prevent_row_split(row)


def _add_field(paragraph, field_code, font_size_pt=9, color_rgb=None):
    """Chen 1 truong (field) Word thuc su (vd PAGE, NUMPAGES) - tu cap nhat dung so trang ke ca
    khi nguoi dung sua/them noi dung vao file Word sau khi xuat, khong phai so ghi cung."""
    run = paragraph.add_run()
    run.font.size = Pt(font_size_pt)
    if color_rgb is not None:
        run.font.color.rgb = color_rgb
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_code} "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    return run


def _add_page_footer(doc, caption_text, logo_bytes=None, show_logo=True):
    """Footer lap lai o MOI trang, giong het dai chan mau (.footer-bar, nen #d9d4c4) + so trang
    "Trang X / Y" (.page-num-label) o goc duoi ban PDF/in - khong con la 1 dong chu tab-stop don
    gian nhu truoc, ma la 1 dai co mau nen that su va so trang tu dong theo Word field."""
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)

    usable_cm = _usable_width_cm(section)

    band = footer.add_table(rows=1, cols=2, width=Cm(usable_cm))
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(band)
    _set_table_col_widths_cm(band, [usable_cm * 0.68, usable_cm * 0.32])
    c0, c1 = band.rows[0].cells
    _shade_cell(c0, FOOTER_BAND_HEX)
    _shade_cell(c1, FOOTER_BAND_HEX)
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(2)
    r1 = p0.add_run(caption_text or "")
    r1.bold = True
    r1.font.color.rgb = FOOTER_CAPTION_RGB
    r1.font.size = Pt(10)

    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo_ok = False
    if show_logo and logo_bytes:
        try:
            p1.add_run().add_picture(BytesIO(logo_bytes), height=Cm(0.6))
            logo_ok = True
        except Exception:
            logo_ok = False
    if not logo_ok and show_logo:
        # Anh logo khong doc duoc (hoac chua co) -> giu chu "ITD Vietnam" nhu mac dinh cu.
        r2 = p1.add_run("ITD")
        r2.bold = True
        r2.font.size = Pt(11)
        r3 = p1.add_run(" Vietnam")
        r3.italic = True
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGBColor(0xA8, 0x82, 0x3C)

    p_num = footer.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_num.paragraph_format.space_before = Pt(2)
    muted = RGBColor(0x6B, 0x66, 0x5A)
    r_lbl = p_num.add_run("Trang ")
    r_lbl.font.size = Pt(9)
    r_lbl.font.color.rgb = muted
    _add_field(p_num, "PAGE", font_size_pt=9, color_rgb=muted)
    r_sep = p_num.add_run(" / ")
    r_sep.font.size = Pt(9)
    r_sep.font.color.rgb = muted
    _add_field(p_num, "NUMPAGES", font_size_pt=9, color_rgb=muted)


def export_word_file(state, save_path):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    # Chong "mo coi" 1 dong bi tach rieng qua trang khac (widow/orphan) - giup van ban chay tu
    # nhien qua cac trang giong nhu ban in/PDF, khong con dong le loi dau/cuoi trang.
    style.paragraph_format.widow_control = True

    for section in doc.sections:
        _set_a4_page(section)

    usable_cm = _usable_width_cm(doc.sections[0])

    show_header_logo = state.get("showHeaderLogo", True)
    show_footer_logo = state.get("showFooterLogo", True)
    logo_bytes = _decode_logo_bytes(state.get("logoDataUri")) if (show_header_logo or show_footer_logo) else None
    footer_caption = _resolve_footer_caption(state)

    _add_page_footer(doc, footer_caption, logo_bytes=logo_bytes, show_logo=show_footer_logo)

    show_vat = bool(state.get("showVatBreakdown"))

    # ---- Header: logo + tieu de ----
    header_table = doc.add_table(rows=1, cols=2)
    _remove_table_borders(header_table)
    _set_table_col_widths_cm(header_table, [usable_cm * 0.45, usable_cm * 0.55])
    c0, c1 = header_table.rows[0].cells
    p0 = c0.paragraphs[0]
    header_logo_ok = False
    if show_header_logo and logo_bytes:
        try:
            p0.add_run().add_picture(BytesIO(logo_bytes), height=Cm(1.2))
            header_logo_ok = True
        except Exception:
            header_logo_ok = False
    if not header_logo_ok and show_header_logo:
        run0 = p0.add_run("ITD")
        run0.bold = True
        run0.font.size = Pt(20)
        run1 = p0.add_run(" Vietnam")
        run1.italic = True
        run1.font.size = Pt(10)
        run1.font.color.rgb = RGBColor(0xA8, 0x82, 0x3C)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_title = p1.add_run("B Á O   G I Á")
    run_title.bold = True
    run_title.font.size = Pt(22)

    p_code = doc.add_paragraph()
    run_code = p_code.add_run(f"Mã số: {state.get('quoteCode','')}")
    run_code.bold = True
    run_code.font.size = Pt(12)

    doc.add_paragraph()

    company_table = doc.add_table(rows=1, cols=2)
    _remove_table_borders(company_table)
    _set_table_col_widths_cm(company_table, [usable_cm * 0.55, usable_cm * 0.45])
    cc0, cc1 = company_table.rows[0].cells
    p_name = cc0.paragraphs[0]
    r_name = p_name.add_run(state.get("companyName", ""))
    r_name.bold = True
    cc0.add_paragraph(state.get("companyAddr", ""))

    p_contact = cc1.paragraphs[0]
    p_contact.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_contact.add_run(f"Tel: {state.get('companyPhone','')}")
    p_contact2 = cc1.add_paragraph()
    p_contact2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_contact2.add_run(f"Email: {state.get('companyEmail','')}")

    doc.add_paragraph()

    p_recipient = doc.add_paragraph()
    r = p_recipient.add_run(f"Kính gửi: {state.get('recipient','')}")
    r.bold = True

    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro_text = f"{state.get('introPrefix','')} {state.get('fixedClause','')}".strip()
    p_intro.add_run(intro_text)
    p_intro.paragraph_format.first_line_indent = Cm(1)

    p_unit = doc.add_paragraph()
    p_unit.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_unit = p_unit.add_run("Đơn vị tính: VNĐ")
    r_unit.italic = True
    r_unit.font.size = Pt(10)

    items = state.get("items", [])
    mode = state.get("mode", "single")
    extra_rows = 0
    if mode == "single":
        if state.get("discountEnabled"):
            extra_rows += 1
        extra_rows += 1

    n_cols = 5 if show_vat else 3
    table = doc.add_table(rows=1 + len(items) + extra_rows, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    hdr_texts = ["TT", "Nội dung", "Đơn giá trước VAT", "VAT", "Đơn giá sau VAT"] if show_vat else ["TT", "Nội dung", "Giá dịch vụ"]
    for i, text in enumerate(hdr_texts):
        hdr[i].paragraphs[0].add_run(text).bold = True
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ty le do rong cot phong theo bang Excel da duyet (TT/Noi dung/gia...) quy doi ve do rong
    # trang A4 thuc te, khong de Word tu chia deu cac cot (nhin se rat khac ban PDF/in).
    col_fracs = [9, 33, 18, 14, 18] if show_vat else [9, 54, 22]
    col_widths_cm = [usable_cm * f / sum(col_fracs) for f in col_fracs]
    _style_data_table(table, col_widths_cm)

    r_idx = 1
    for it in items:
        is_main = bool(it.get("isMain"))
        # Dong hang muc chinh co the duoc bat tuy chon "cong tong cac hang muc con vao dong nay"
        # (xem g.sumSubItems/isSumRow ben ui.html) - khi do no hien gia tri (in dam) giong 1 dong
        # binh thuong thay vi merge/de trong nhu mac dinh.
        is_sum_row = is_main and bool(it.get("isSumRow"))
        cells = table.rows[r_idx].cells
        stt_run = cells[0].paragraphs[0].add_run(str(it.get("number", "")))
        stt_run.bold = is_main
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = cells[1].paragraphs[0].add_run(it.get("name", ""))
        name_run.bold = is_main
        if is_main and not is_sum_row:
            if show_vat:
                cells[1].merge(cells[n_cols - 1])
        else:
            # Dong hang muc con: cho phep in nghieng toan dong neu nguoi dung bat tuy chon (xem
            # g.subItemsItalic/row.italic ben ui.html) - mac dinh chu thang nhu truoc gio.
            is_italic = (not is_main) and bool(it.get("italic"))
            if is_italic:
                stt_run.italic = True
                name_run.italic = True
            if show_vat:
                bv = it.get("beforeVat", None)
                va = it.get("vatAmount", None)
                if bv is not None:
                    p2 = cells[2].paragraphs[0]
                    r2 = p2.add_run(fmt_number(bv))
                    r2.bold = is_main
                    r2.italic = is_italic
                    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if va is not None:
                    p3 = cells[3].paragraphs[0]
                    r3 = p3.add_run(fmt_number(va))
                    r3.bold = is_main
                    r3.italic = is_italic
                    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            amt = it.get("amount", None)
            if amt is not None:
                p_last = cells[n_cols - 1].paragraphs[0]
                r_last = p_last.add_run(fmt_number(amt))
                r_last.bold = is_main
                r_last.italic = is_italic
                p_last.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_idx += 1

    if mode == "single":
        if state.get("discountEnabled"):
            cells = table.rows[r_idx].cells
            pct = state.get("discountPct", 0)
            amt = -round(state.get("discountAmount", 0))
            discount_label = "Giảm giá dịch vụ" if state.get("discountMode") == "fixed" else f"Giảm giá dịch vụ {pct}%"
            r_disc_label = cells[1].paragraphs[0].add_run(discount_label)
            r_disc_label.italic = True
            r_disc_label.font.color.rgb = ACCENT_RGB
            if show_vat:
                cells[1].merge(cells[n_cols - 2])
            p_last = cells[n_cols - 1].paragraphs[0]
            r_disc_amt = p_last.add_run(fmt_number(amt))
            r_disc_amt.italic = True
            r_disc_amt.font.color.rgb = ACCENT_RGB
            p_last.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_idx += 1
        cells = table.rows[r_idx].cells
        cells[1].paragraphs[0].add_run("Tổng cộng giá trị dịch vụ").bold = True
        if show_vat:
            cells[1].merge(cells[n_cols - 2])
        p_last = cells[n_cols - 1].paragraphs[0]
        run_total = p_last.add_run(fmt_number(round(state.get("grandTotal", 0))))
        run_total.bold = True
        p_last.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        # Hang "Tong cong" to mau nhe giong .quote-table tr.total-row trong ban PDF/in.
        _shade_row(table.rows[r_idx], MAIN_TOTAL_ROW_HEX)
        r_idx += 1

    _set_table_font_size(table, 11)

    doc.add_paragraph()

    for note in state.get("footerNotes", []):
        p = doc.add_paragraph()
        p.add_run(note)

    p_closing = doc.add_paragraph()
    p_closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_closing = p_closing.add_run(state.get("closingNote", ""))
    r_closing.bold = True

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.add_run(
        f"Hà Nội, ngày {state.get('qDay','')} tháng {state.get('qMonth','')} năm {state.get('qYear','')}"
    )

    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sign_company_name = (state.get("signCompanyName") or state.get("companyName", "") or "").upper()
    r_sign = p_sign.add_run(sign_company_name)
    r_sign.bold = True

    # ---- Cac trang Du toan chi tiet theo hang muc ----
    salary_table = state.get("salaryTable", {"levels": [], "workDaysPerMonth": 26})
    levels_by_id = {lvl["id"]: lvl for lvl in salary_table.get("levels", [])}

    cql_tiers = state.get("cqlTiers")
    for detail in state.get("detailSheets", []):
        doc.add_page_break()
        totals = compute_detail_totals(detail, salary_table, cql_tiers)
        daily_rate = totals["daily_rate"]

        h = doc.add_paragraph()
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.keep_with_next = True
        rh = h.add_run(f"TỔNG HỢP CHI PHÍ: {detail.get('name','').upper()}")
        rh.bold = True
        rh.font.size = Pt(15)

        round_label = "TỔNG CỘNG (LÀM TRÒN)" if totals["roundEnabled"] else "TỔNG CỘNG"

        ccg_explain = "Trọn gói" if detail.get("ccgMode") == "lumpsum" else "Tổng hợp chi phí chuyên gia"
        ck_explain = "Trọn gói" if detail.get("ckMode") == "lumpsum" else "Tổng hợp chi phí khác"
        rows_data = [
            ("I", "CHI PHÍ CHUYÊN GIA", "Ccg", ccg_explain, totals["Ccg"], False),
            ("II", "CHI PHÍ QUẢN LÝ", "Cql", f"({totals['cqlPct']}%)*Ccg", totals["Cql"], False),
            ("III", "CHI PHÍ KHÁC", "Ck", ck_explain, totals["Ck"], False),
            ("IV", "THU NHẬP CHỊU THUẾ TÍNH TRƯỚC", "TN", f"{detail.get('tnPct',0)}%*(Ccg+Cql+Ck)", totals["TN"], False),
            ("", "CỘNG", "G", "Ccg+Cql+Ck+TN", totals["G"], True),
        ]
        if show_vat:
            # Bang bao gia chinh dang hien thi chi tiet VAT: giam gia tru NGAY tren G (truoc thue),
            # roi moi tinh VAT tren phan da giam - dong bo voi cach tach truoc VAT/VAT o bang chinh.
            # Neu la giam gia so tien co dinh: khong ghi noi dung vao cot "Diễn giải" (de trong).
            # LUON them 1 dong "lam tron" gia tri TRUOC THUE (mainVatBase) truoc dong Thue - ca khi CO
            # va KHONG CO giam gia - vi dong Thue (mainVatAmount) va TONG CONG (roundedTotal) o duoi deu
            # da tinh tren cac so DA lam tron; neu thieu dong nay, CỘNG (tho) + Thue (da lam tron) se
            # khong con khop voi TONG CONG, va cung khong khop voi cot "Truoc VAT" cua bang bao gia
            # chinh (cung dung dung mainVatBase) - dung yeu cau nguoi dung da neu ve viec khop hoan toan.
            discount_desc = "" if detail.get("discountMode") == "fixed" else f"{detail.get('discountPct',0)}%"
            round_before_tax_label = "GIÁ TRỊ TRƯỚC THUẾ (LÀM TRÒN)" if totals["roundEnabled"] else "GIÁ TRỊ TRƯỚC THUẾ"
            if detail.get("discountEnabled"):
                round_before_tax_label_disc = "GIÁ TRỊ TRƯỚC THUẾ (SAU GIẢM GIÁ) (LÀM TRÒN)" if totals["roundEnabled"] else "GIÁ TRỊ TRƯỚC THUẾ (SAU GIẢM GIÁ)"
                rows_data.append(("", "Giảm giá dịch vụ", "", discount_desc, -(totals["G"] - totals["mainVatBase"]), False))
                rows_data.append(("", round_before_tax_label_disc, "G'", "G - Giảm giá dịch vụ", totals["mainVatBase"], True))
                vat_explain = f"{detail.get('vatPct',0)}% x G'"
            else:
                rows_data.append(("", round_before_tax_label, "G", "(làm tròn)", totals["mainVatBase"], True))
                vat_explain = f"{detail.get('vatPct',0)}% x G"
            rows_data.append(("V", "Thuế giá trị gia tăng", "GTGT", vat_explain, totals["mainVatAmount"], False))
        else:
            rows_data.append(("V", "Thuế giá trị gia tăng", "GTGT", f"{detail.get('vatPct',0)}% x G", totals["VAT"], False))
            rows_data.append(("", "TỔNG CỘNG", "", "G + GTGT", totals["beforeDiscount"], True))
            if detail.get("discountEnabled"):
                discount_desc = "" if detail.get("discountMode") == "fixed" else f"{detail.get('discountPct',0)}%"
                rows_data.append(("", "Giảm giá dịch vụ", "", discount_desc, -totals["discount"], False))
        rows_data.append(("", round_label, "", "", totals["roundedTotal"], True))

        # Bang Tong hop chi phi
        t0 = doc.add_table(rows=1 + len(rows_data), cols=5)
        t0.style = "Table Grid"
        hdr0 = t0.rows[0].cells
        for i, text in enumerate(["STT", "Khoản mục chi phí", "Ký hiệu", "Diễn giải", "Thành tiền"]):
            hdr0[i].paragraphs[0].add_run(text).bold = True
            hdr0[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_data_table(t0, [usable_cm * f / 100 for f in (7, 25, 10, 34, 24)])

        for ridx, (stt, label, sym, explain, val, is_bold) in enumerate(rows_data, start=1):
            cells = t0.rows[ridx].cells
            cells[0].paragraphs[0].add_run(stt).bold = is_bold
            cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cells[1].paragraphs[0].add_run(label).bold = is_bold
            cells[2].paragraphs[0].add_run(sym).bold = is_bold
            cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cells[3].paragraphs[0].add_run(explain).bold = is_bold
            p4 = cells[4].paragraphs[0]
            p4.add_run(fmt_number(val)).bold = is_bold
            p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            # Cac dong "CONG"/"TONG CONG..." to mau nhe giong tr.sum-row trong ban PDF/in.
            if is_bold:
                _shade_row(t0.rows[ridx], SUM_ROW_HEX)

        _set_table_font_size(t0, 10.5)

        # Dien giai chi phi chuyen gia
        if detail.get("ccgMode") != "lumpsum":
            doc.add_paragraph()
            h2 = doc.add_paragraph()
            h2.paragraph_format.keep_with_next = True
            h2.add_run("Diễn giải chi phí chuyên gia").bold = True

            expert_rows = detail.get("expertRows", [])
            analysis = totals["expert_analysis"]
            # Moi muc chuyen gia trong 1 cong viec xuong 1 dong rieng -> tinh truoc tong so
            # dong can thiet (2 dong dau la header + dong tong hop) de tao bang dung kich thuoc.
            row_counts = []
            for i, er in enumerate(expert_rows):
                if analysis[i]["hasChildren"]:
                    row_counts.append(1)
                else:
                    row_levels = er.get("levels") or [{
                        "levelId": er.get("levelId"), "quantity": er.get("quantity"), "days": er.get("days"),
                    }]
                    row_counts.append(max(1, len(row_levels)))
            t1 = doc.add_table(rows=2 + sum(row_counts), cols=7)
            t1.style = "Table Grid"
            hdr1 = t1.rows[0].cells
            for i, text in enumerate(
                ["STT", "Công việc thực hiện", "Mức chuyên gia", "Khối lượng chuyên gia",
                 "Ngày công", "Đơn giá ngày công", "Thành tiền (VNĐ)"]
            ):
                hdr1[i].paragraphs[0].add_run(text).bold = True
                hdr1[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_data_table(t1, [usable_cm * f / 100 for f in (7, 29, 12, 12, 10, 15, 15)])

            sum_cells = t1.rows[1].cells
            sum_cells[1].paragraphs[0].add_run("Tổng hợp chi phí chuyên gia tư vấn").bold = True
            p6 = sum_cells[6].paragraphs[0]
            p6.add_run(fmt_number(totals["Ccg"])).bold = True
            p6.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _shade_row(t1.rows[1], SUM_ROW_HEX)

            expert_nums = compute_expert_row_numbers(expert_rows)
            ridx = 2
            for i, er in enumerate(expert_rows):
                a = analysis[i]
                depth = a.get("depth", 0) or 0
                if a["hasChildren"]:
                    cells = t1.rows[ridx].cells
                    cells[0].paragraphs[0].add_run(expert_nums[i])
                    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p1_desc = cells[1].paragraphs[0]
                    p1_desc.add_run(er.get("desc", ""))
                    if depth:
                        p1_desc.paragraph_format.left_indent = Cm(0.4 * depth)
                    for ci in (2, 3, 4, 5):
                        cells[ci].paragraphs[0].add_run("-")
                        cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p6r = cells[6].paragraphs[0]
                    r6 = p6r.add_run(fmt_number(a["childrenTotal"]))
                    r6.bold = True
                    p6r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    ridx += 1
                else:
                    # Moi muc chuyen gia -> 1 dong; cot STT/Cong viec thuc hien gop (merge)
                    # doc theo ca nhom cac dong nay (giong bang mau Excel da duoc duyet).
                    row_levels = er.get("levels") or [{
                        "levelId": er.get("levelId"), "quantity": er.get("quantity"), "days": er.get("days"),
                    }]
                    group_start = ridx
                    for lv in row_levels:
                        lvl = levels_by_id.get(lv.get("levelId"))
                        dr = daily_rate(lv.get("levelId"))
                        qty = lv.get("quantity") or 0
                        days = lv.get("days") or 0
                        cells = t1.rows[ridx].cells
                        p2 = cells[2].paragraphs[0]
                        p2.add_run(lvl.get("name") if lvl else "")
                        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p3 = cells[3].paragraphs[0]
                        p3.add_run(str(qty))
                        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p4 = cells[4].paragraphs[0]
                        p4.add_run(str(days))
                        p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p5 = cells[5].paragraphs[0]
                        p5.add_run(fmt_number(dr))
                        p5.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        p6c = cells[6].paragraphs[0]
                        p6c.add_run(fmt_number(qty * days * dr))
                        p6c.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        ridx += 1
                    group_end = ridx - 1
                    first_cells = t1.rows[group_start].cells
                    first_cells[0].paragraphs[0].add_run(expert_nums[i])
                    first_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p1_leaf_desc = first_cells[1].paragraphs[0]
                    p1_leaf_desc.add_run(er.get("desc", ""))
                    if depth:
                        p1_leaf_desc.paragraph_format.left_indent = Cm(0.4 * depth)
                    if group_end > group_start:
                        t1.cell(group_start, 0).merge(t1.cell(group_end, 0))
                        t1.cell(group_start, 1).merge(t1.cell(group_end, 1))
                        t1.cell(group_start, 0).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        t1.cell(group_start, 1).vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            _set_table_font_size(t1, 10)

        # Dien giai chi phi khac
        other_rows = detail.get("otherRows", [])
        if detail.get("ckMode") != "lumpsum" and any(r.get("desc") or r.get("quantity") or r.get("unitPrice") for r in other_rows):
            doc.add_paragraph()
            h3 = doc.add_paragraph()
            h3.paragraph_format.keep_with_next = True
            h3.add_run("Diễn giải chi phí khác").bold = True

            t2 = doc.add_table(rows=2 + len(other_rows), cols=6)
            t2.style = "Table Grid"
            hdr2 = t2.rows[0].cells
            for i, text in enumerate(["STT", "Vật tư, vật liệu", "Đơn vị", "Khối lượng", "Đơn giá (đ)", "Thành tiền (đ)"]):
                hdr2[i].paragraphs[0].add_run(text).bold = True
                hdr2[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _style_data_table(t2, [usable_cm * f / 100 for f in (7, 33, 12, 14, 17, 17)])

            sum_cells2 = t2.rows[1].cells
            sum_cells2[1].paragraphs[0].add_run("Tổng cộng").bold = True
            p5b = sum_cells2[5].paragraphs[0]
            p5b.add_run(fmt_number(totals["Ck"])).bold = True
            p5b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _shade_row(t2.rows[1], SUM_ROW_HEX)

            other_nums = compute_row_numbers(other_rows)
            for i, orow in enumerate(other_rows):
                cells = t2.rows[2 + i].cells
                qty = orow.get("quantity") or 0
                price = orow.get("unitPrice") or 0
                cells[0].paragraphs[0].add_run(other_nums[i])
                cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cells[1].paragraphs[0].add_run(orow.get("desc", ""))
                cells[2].paragraphs[0].add_run(orow.get("unit", ""))
                cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cells[3].paragraphs[0].add_run(str(qty))
                cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                p4b = cells[4].paragraphs[0]
                p4b.add_run(fmt_number(price))
                p4b.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p5c = cells[5].paragraphs[0]
                p5c.add_run(fmt_number(qty * price))
                p5c.alignment = WD_ALIGN_PARAGRAPH.RIGHT

            _set_table_font_size(t2, 10.5)

    # ---- Trang Bang luong chuyen gia (dung chung) ----
    # CHI xuat trang nay khi co IT NHAT 1 hang muc dang dung ccgMode "linked" (Link tu dien giai) - tuc
    # co thuc su tham chieu den bang luong. Neu TAT CA hang muc deu chon "Tron goi" (lumpsum) thi bang
    # luong khong duoc dung den o dau ca, xuat ra se thua.
    any_sheet_uses_salary_table = any(d.get("ccgMode") == "linked" for d in state.get("detailSheets", []))
    if any_sheet_uses_salary_table:
        doc.add_page_break()
        h_sal = doc.add_paragraph()
        h_sal.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h_sal.paragraph_format.keep_with_next = True
        rh_sal = h_sal.add_run("BẢNG LƯƠNG CHUYÊN GIA")
        rh_sal.bold = True
        rh_sal.font.size = Pt(15)

        p_sal_note = doc.add_paragraph()
        p_sal_note.paragraph_format.keep_with_next = True
        days_per_month = salary_table.get("workDaysPerMonth", 26) or 26
        p_sal_note.add_run(
            f"Áp dụng theo: {salary_table.get('circularName','')} — quy đổi {days_per_month} ngày công/tháng"
        )

        levels = salary_table.get("levels", [])
        t_sal = doc.add_table(rows=1 + len(levels), cols=3)
        t_sal.style = "Table Grid"
        hdr_sal = t_sal.rows[0].cells
        for i, text in enumerate(["Mức", "Định mức/tháng", "Định mức/ngày"]):
            hdr_sal[i].paragraphs[0].add_run(text).bold = True
            hdr_sal[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_data_table(t_sal, [usable_cm * f / 100 for f in (30, 35, 35)])
        for i, lvl in enumerate(levels, start=1):
            cells = t_sal.rows[i].cells
            cells[0].paragraphs[0].add_run(lvl.get("name", ""))
            p2 = cells[1].paragraphs[0]
            p2.add_run(fmt_number(lvl.get("monthlyRate", 0)))
            p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p3 = cells[2].paragraphs[0]
            p3.add_run(fmt_number(lvl.get("monthlyRate", 0) / days_per_month))
            p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        _set_table_font_size(t_sal, 11)

    doc.save(save_path)


# =====================================================================================
# ---- Import / Export file mau Excel (chuyen gia / chi phi khac / bang luong) ----
# Muc dich: cho nguoi dung dien du lieu vao 1 file Excel theo dung chuan cot cua phan
# mem, roi Import lai thay vi phai nhap tay tung dong tren giao dien.
# =====================================================================================

def _strip_diacritics(s):
    """Bo dau tieng Viet de so sanh chuoi khong phu thuoc co dau/khong dau (nguoi dung
    co the go tieu de co dau hoac khong dau deu nhan dien duoc)."""
    s = str(s)
    s = s.replace("đ", "d").replace("Đ", "D")
    nfkd = unicodedata.normalize("NFD", s)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def _find_header_row(ws, marker_text, max_scan=15):
    """Do tim dong tieu de trong ws: dong dau tien co 1 o chua chuoi marker_text
    (khong phan biet hoa/thuong, khong phan biet co dau/khong dau tieng Viet).
    Tra ve so thu tu dong (1-based) hoac None neu khong thay."""
    marker_norm = _strip_diacritics(marker_text).strip().lower()
    for r in range(1, min(max_scan, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v and marker_norm in _strip_diacritics(v).strip().lower():
                return r
    return None


def _template_base_styles():
    return {
        "bold": Font(bold=True, name="Times New Roman", size=12),
        "normal": Font(name="Times New Roman", size=11),
        "header_fill": PatternFill("solid", fgColor="D9D2C0"),
        "border": Border(
            left=Side(style="thin", color="999999"), right=Side(style="thin", color="999999"),
            top=Side(style="thin", color="999999"), bottom=Side(style="thin", color="999999"),
        ),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
    }


def _write_instructions_sheet(wb, title, lines):
    hd = wb.create_sheet("Huong dan")
    hd.column_dimensions["A"].width = 100
    hd.cell(row=1, column=1, value=title).font = Font(bold=True, name="Times New Roman", size=13, color="8A3B2D")
    for i, line in enumerate(lines, start=2):
        hd.cell(row=i, column=1, value=line).font = Font(name="Times New Roman", size=12)


# ---- 1. Danh sach cong viec chuyen gia (expertRows) ----

def build_expert_template_workbook(salary_table):
    """Tao file Excel mau de dien danh sach cong viec chuyen gia. Cot 'Muc chuyen gia'
    co dropdown lay dung theo ten cac muc dang co trong Bang luong chuyen gia hien tai
    cua bao gia, de tranh go sai ten khi import lai."""
    st = _template_base_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Du lieu"

    ws.merge_cells("A1:F1")
    title_cell = ws.cell(
        row=1, column=1,
        value="FILE MẪU — Danh sách công việc chuyên gia (điền theo đúng cấu trúc này rồi Import lại vào phần mềm)",
    )
    title_cell.font = Font(bold=True, size=12, name="Times New Roman", color="8A3B2D")
    title_cell.alignment = st["left"]

    headers = ["STT", "Công việc thực hiện", "Mức chuyên gia", "Số người", "Ngày công", "Cấp (1=cao nhất, 2=con, 3=cháu...)"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = st["bold"]
        c.fill = st["header_fill"]
        c.alignment = st["center"]
        c.border = st["border"]

    level_names = [lvl.get("name", "") for lvl in salary_table.get("levels", [])] or ["Mức 1"]

    # Vi du minh hoa cau truc NHIEU CAP (khong gioi han o 2 cap me/con nhu truoc) - mo phong dung
    # vi du cua nguoi dung: "3 Thuc hien khao sat" (cap 1) -> "3.1 Tai Cuc Quan ly no..." (cap 2)
    # -> "3.1.1/3.1.2 Hang muc phan cung/phan mem" (cap 3, 2 dong la). Dong cap 1 va cap 2 deu la
    # dong "me" (co dong con ngay sau) nen de TRONG so lieu rieng (Muc/So nguoi/Ngay cong) - giong
    # dung cach giao dien tu khoa cac o nay khi 1 dong co dong con; chi dong LA (khong co dong con
    # nao sau no) moi can dien so lieu thuc.
    example_rows = [
        (1, "Ví dụ: Thực hiện khảo sát (dòng mẹ cấp 1 — để trống số liệu vì có dòng con)", "", "", "", 1),
        (None, "Ví dụ: Tại Cục Quản lý nợ và Kinh tế đối ngoại (con của dòng trên, cấp 2 — cũng để trống vì có dòng con riêng)", "", "", "", 2),
        (None, "Ví dụ: Hạng mục phần cứng máy tính (dòng lá cấp 3 — điền số liệu thật ở đây)", level_names[0], 1, 5, 3),
        (None, "Ví dụ: Hạng mục phần mềm nội bộ (dòng lá cấp 3 khác, cùng là con của dòng cấp 2 trên)", level_names[0], 1, 20, 3),
        (2, "Ví dụ: Công việc độc lập khác (dòng cấp 1 mới, không phải con của ai)", level_names[min(1, len(level_names) - 1)], 2, 5, 1),
    ]
    start_row = 3
    for offset, (stt, desc, lvl_name, qty, days, cap) in enumerate(example_rows):
        r = start_row + offset
        ws.cell(row=r, column=1, value=stt).alignment = st["center"]
        ws.cell(row=r, column=2, value=desc).alignment = st["left"]
        ws.cell(row=r, column=3, value=lvl_name).alignment = st["center"]
        ws.cell(row=r, column=4, value=qty).alignment = st["center"]
        ws.cell(row=r, column=5, value=days).alignment = st["center"]
        ws.cell(row=r, column=6, value=cap).alignment = st["center"]
        for c in range(1, 7):
            ws.cell(row=r, column=c).font = st["normal"]
            ws.cell(row=r, column=c).border = st["border"]

    last_example_row = start_row + len(example_rows) - 1
    blank_rows_end = last_example_row + 30
    for r in range(last_example_row + 1, blank_rows_end + 1):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = st["border"]
            ws.cell(row=r, column=c).font = st["normal"]

    dv_formula = '"' + ",".join(level_names) + '"'
    if len(dv_formula) <= 255:
        dv = DataValidation(type="list", formula1=dv_formula, allow_blank=True, showDropDown=False)
        dv.error = "Vui long chon 1 trong cac muc chuyen gia da co trong Bang luong chuyen gia."
        dv.errorTitle = "Muc chuyen gia khong hop le"
        ws.add_data_validation(dv)
        dv.add(f"C{start_row}:C{blank_rows_end}")

    dv_level = DataValidation(
        type="whole", operator="between", formula1=1, formula2=20, allow_blank=True, showErrorMessage=True,
    )
    dv_level.error = "Cột 'Cấp' phải là số nguyên từ 1 trở lên (1 = cấp cao nhất, 2 = con, 3 = cháu...)."
    dv_level.errorTitle = "Cấp không hợp lệ"
    ws.add_data_validation(dv_level)
    dv_level.add(f"F{start_row}:F{blank_rows_end}")

    # Canh bao (khong chan cung, chi hien popup co the bo qua) o cot Muc chuyen gia/So nguoi/Ngay cong:
    # neu dong NGAY BEN DUOI co Cap SO LON HON dong hien tai (tuc la dong duoi la CON cua dong nay), thi
    # dong hien tai la dong "me" - so lieu rieng vua nhap se KHONG duoc phan mem su dung vi tong se tu
    # dong lay theo (cac) dong con, giong dung cach giao dien tu khoa cac o nay khi phat hien co dong con.
    # Dung 1 formula tuong doi (chi khoa cot F bang $), ap dung cho toan bo khoi C{start}:E{end} — Excel
    # tu dich chuyen tham chieu dong cho tung o giong cong thuc thuong, nen luon kiem tra dung dong ngay
    # ben duoi cua moi dong du lieu.
    dv_parent_hint = DataValidation(
        type="custom",
        formula1=f"$F{start_row + 1}<=$F{start_row}",
        allow_blank=True,
        showErrorMessage=True,
    )
    dv_parent_hint.errorStyle = "warning"
    dv_parent_hint.error = (
        "Dòng NGAY BÊN DƯỚI đang có số Cấp LỚN HƠN dòng này (tức là dòng dưới là CON của dòng này).\n"
        "Nếu đúng vậy: số liệu Mức chuyên gia/Số người/Ngày công anh vừa nhập ở đây sẽ KHÔNG được phần mềm "
        "sử dụng khi import — tổng của dòng này sẽ tự động lấy bằng tổng các dòng con.\n"
        "Bấm 'Có' để vẫn giữ giá trị đã nhập (ví dụ nếu dòng dưới thực ra không phải con), hoặc 'Không' để xoá lại."
    )
    dv_parent_hint.errorTitle = "Có thể là dòng mẹ — số liệu này sẽ bị bỏ qua"
    ws.add_data_validation(dv_parent_hint)
    dv_parent_hint.add(f"C{start_row}:E{blank_rows_end}")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 26
    ws.freeze_panes = "A3"

    _write_instructions_sheet(wb, "HƯỚNG DẪN DIỀN FILE MẪU — Danh sách công việc chuyên gia", [
        "",
        "1. Mỗi dòng là 1 'Mức chuyên gia' tham gia 1 công việc.",
        "2. Nếu 1 công việc có NHIỀU mức chuyên gia cùng tham gia: điền cột 'Công việc thực hiện' ở dòng ĐẦU TIÊN,",
        "   các dòng TIẾP THEO của cùng công việc đó CHỈ CẦN để trống cột STT và Công việc thực hiện,",
        "   chỉ điền tiếp cột Mức chuyên gia / Số người / Ngày công.",
        "3. Cột 'Mức chuyên gia' PHẢI ghi ĐÚNG TÊN đang có trong Bảng lương chuyên gia của báo giá",
        "   (có sẵn dropdown để chọn, tránh gõ sai chính tả). Nếu tên không khớp, khi import phần mềm sẽ cảnh báo",
        "   và tạm dùng một mức mặc định — anh nên sửa lại cho đúng trước khi import.",
        "4. Cột 'Số người' và 'Ngày công' là số, không điền chữ.",
        "5. CỘT 'CẤP' — hỗ trợ NHIỀU CẤP lồng nhau (không chỉ 2 cấp mẹ/con như trước): điền số 1 cho công việc",
        "   ở cấp cao nhất, số 2 cho công việc là CON của 1 công việc cấp 1 GẦN NHẤT phía trên nó, số 3 cho",
        "   công việc là CON của 1 công việc cấp 2 gần nhất phía trên, và tiếp tục như vậy. CHỈ được tăng tối đa",
        "   1 cấp so với dòng NGAY PHÍA TRÊN (ví dụ dòng trên là cấp 2 thì dòng này tối đa là cấp 3, không được",
        "   nhảy thẳng lên cấp 4) — nếu lỡ điền vượt quá, phần mềm sẽ tự điều chỉnh về đúng cấp tối đa cho phép.",
        "   Để TRỐNG cột này thì phần mềm sẽ tự hiểu là CÙNG CẤP với dòng công việc ngay phía trên.",
        "6. DÒNG CÓ DÒNG CON (ví dụ dòng cấp 1 mà dòng ngay sau là cấp 2): tổng tiền của dòng này sẽ TỰ ĐỘNG",
        "   bằng tổng các dòng con trực tiếp của nó — phần mềm bỏ qua số liệu Mức/Số người/Ngày công đã nhập ở",
        "   dòng đó (nên cứ để trống cho đỡ nhầm, xem 2 dòng ví dụ đầu). CHỈ dòng 'lá' (không có dòng con nào)",
        "   mới cần điền số liệu thật — xem 2 dòng ví dụ cấp 3 ở trên.",
        "7. LƯU Ý khi gõ Mức chuyên gia/Số người/Ngày công: nếu dòng NGAY BÊN DƯỚI dòng anh đang gõ có Cấp lớn",
        "   hơn (là dòng con), Excel sẽ hiện popup cảnh báo dòng anh đang gõ có thể là dòng MẸ và số liệu đó sẽ",
        "   không được phần mềm sử dụng. Đây chỉ là lời nhắc (không bắt buộc) — nếu đúng ý anh vẫn bấm 'Có' để",
        "   giữ lại giá trị đã gõ.",
        "8. Sau khi điền xong, lưu file lại (giữ nguyên định dạng .xlsx), vào phần mềm bấm nút",
        "   'Nhập từ Excel' ở phần 'Chi phí chuyên gia' của hạng mục cần nhập, chọn đúng file này.",
        "9. Import sẽ THAY THẾ toàn bộ danh sách công việc chuyên gia hiện tại của hạng mục đó — anh nên kiểm tra",
        "   lại sau khi import.",
    ])

    return wb


_SUB_FLAG_TRUE_VALUES = {"x", "co", "yes", "y", "true", "1"}


def _is_sub_flag_true(raw_value):
    """Dung cho otherRows (parse_other_costs_excel) - otherRows van giu nguyen mo hinh 2 cap
    me/con cu (isSub boolean), KHONG doi sang Cap nhieu muc nhu expertRows."""
    if raw_value is None:
        return False
    return _strip_diacritics(str(raw_value)).strip().lower() in _SUB_FLAG_TRUE_VALUES


def _parse_expert_cap_value(raw_value):
    """Doc 1 gia tri o cot 'Cap' - tra ve so nguyen (>=1) neu doc duoc, hoac None neu de trong/khong
    doc duoc (dong bo voi quy uoc: de trong = cung cap voi dong truoc)."""
    if raw_value is None or str(raw_value).strip() == "":
        return None
    try:
        return max(1, int(round(float(raw_value))))
    except (TypeError, ValueError):
        return None


def parse_expert_import_excel(path, salary_table):
    """Doc file Excel da dien theo file mau va chuyen thanh danh sach expertRows:
    [{desc, depth, levels:[{levelId, quantity, days}, ...]}, ...].
    Cac dong lien tiep co cot 'Cong viec' de trong duoc gop vao 'levels' cua dong cong
    viec gan nhat (dung 1 dong = 1 cong viec, nhieu muc chuyen gia trong cung dong).
    Cot 'Cap' (cot F) la so NGUYEN tu 1 tro len, cho phep NHIEU CAP long nhau (khong gioi
    han o 2 cap me/con nhu truoc): 1 = cap cao nhat, 2 = con cua 1 dong cap 1 gan nhat phia
    truoc, 3 = con cua 1 dong cap 2 gan nhat phia truoc... Moi dong chi duoc tang toi da 1
    cap so voi dong NGAY PHIA TRUOC (khop voi quy tac maxAllowedExpertDepth() trong ui.html) -
    vuot qua se tu dong bi gioi han lai va ghi vao warnings. De trong cot Cap = hieu la CUNG
    CAP voi dong cong viec ngay phia truoc. Khop 'Muc chuyen gia' voi salary_table['levels']
    theo ten (bo qua hoa/thuong, khoang trang thua); ten khong khop se dung tam muc dau tien
    va ghi vao warnings. Tra ve (rows, warnings)."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Du lieu"] if "Du lieu" in wb.sheetnames else wb.worksheets[0]

    header_row = _find_header_row(ws, "muc chuyen gia") or _find_header_row(ws, "cong viec")
    if header_row is None:
        raise ValueError("Không tìm thấy dòng tiêu đề trong file — anh dùng đúng file mẫu nhé.")

    levels = salary_table.get("levels", [])
    levels_by_name = {(lvl.get("name") or "").strip().lower(): lvl["id"] for lvl in levels if lvl.get("name")}
    fallback_level_id = levels[0]["id"] if levels else None

    rows = []
    warnings = []
    current = None
    prev_depth = -1  # depth (0-based) cua dong cong viec truoc do; -1 = chua co dong nao

    for r in range(header_row + 1, ws.max_row + 1):
        stt = ws.cell(row=r, column=1).value
        desc = ws.cell(row=r, column=2).value
        level_name = ws.cell(row=r, column=3).value
        qty = ws.cell(row=r, column=4).value
        days = ws.cell(row=r, column=5).value
        cap_raw = ws.cell(row=r, column=6).value

        if all(v is None or str(v).strip() == "" for v in [stt, desc, level_name, qty, days, cap_raw]):
            continue

        desc_text = str(desc).strip() if desc is not None else ""
        level_text = str(level_name).strip() if level_name is not None else ""

        if desc_text:
            cap_val = _parse_expert_cap_value(cap_raw)
            max_allowed_depth = 0 if prev_depth < 0 else prev_depth + 1
            if cap_val is None:
                depth = 0 if prev_depth < 0 else prev_depth
            else:
                depth = cap_val - 1
                if depth > max_allowed_depth:
                    warnings.append(
                        f"Dòng {r}: Cấp {depth + 1} vượt quá cấp tối đa cho phép ở vị trí này "
                        f"(chỉ được tăng tối đa 1 cấp so với dòng ngay phía trên) — đã tự điều chỉnh về Cấp {max_allowed_depth + 1}."
                    )
                    depth = max_allowed_depth
                elif depth < 0:
                    depth = 0
            prev_depth = depth
            current = {"desc": desc_text, "depth": depth, "levels": []}
            rows.append(current)
        elif current is None:
            warnings.append(f"Dòng {r}: bỏ qua vì thiếu cột 'Công việc thực hiện' và chưa có dòng công việc nào trước đó.")
            continue

        if not level_text:
            continue

        level_id = levels_by_name.get(level_text.lower())
        if not level_id:
            warnings.append(f"Dòng {r}: không tìm thấy mức chuyên gia '{level_text}' trong bảng lương hiện tại — đã dùng tạm mức đầu tiên, anh kiểm tra lại.")
            level_id = fallback_level_id

        try:
            qty_val = float(qty) if qty is not None and str(qty).strip() != "" else 1
        except (TypeError, ValueError):
            qty_val = 1
        try:
            days_val = float(days) if days is not None and str(days).strip() != "" else 0
        except (TypeError, ValueError):
            days_val = 0

        if level_id:
            current["levels"].append({"levelId": level_id, "quantity": qty_val, "days": days_val})

    # Dong co dong con (dong ngay sau no sau trong hon) thi KHONG can "levels" rieng - so lieu cua no
    # se bi bo qua hoan toan boi analyze_expert_rows() (tu dong lay tong cac dong con), nen KHONG dien
    # tam 1 muc mac dinh/canh bao cho nhung dong nay (vua thua, vua de gay nham lan vi file mau co huong
    # dan de trong cac dong me).
    n_rows = len(rows)
    cleaned = []
    for i, row in enumerate(rows):
        has_children = i + 1 < n_rows and rows[i + 1]["depth"] > row["depth"]
        if not row["levels"] and not has_children:
            if fallback_level_id:
                row["levels"] = [{"levelId": fallback_level_id, "quantity": 1, "days": 0}]
                warnings.append(f"Công việc '{row['desc']}' không có mức chuyên gia hợp lệ nào — đã điền tạm 1 mức mặc định, anh nhập lại số liệu.")
            else:
                continue
        cleaned.append(row)

    if not cleaned:
        raise ValueError("Không đọc được dòng dữ liệu nào hợp lệ trong file.")

    return cleaned, warnings


# ---- 2. Danh sach chi phi khac (otherRows) ----

def build_other_costs_template():
    st = _template_base_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Du lieu"

    ws.merge_cells("A1:E1")
    title_cell = ws.cell(
        row=1, column=1,
        value="FILE MAU — Danh sach chi phi khac (dien theo dung cau truc nay roi Import lai vao phan mem)",
    )
    title_cell.font = Font(bold=True, size=12, name="Times New Roman", color="8A3B2D")
    title_cell.alignment = st["left"]

    headers = ["Mô tả (vật tư, vật liệu...)", "Đơn vị tính", "Khối lượng", "Đơn giá", "Là dòng con? (x = có)"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = st["bold"]
        c.fill = st["header_fill"]
        c.alignment = st["center"]
        c.border = st["border"]

    # Vi du minh hoa hang muc me - hang muc con: dong 2 danh dau "x" o cot cuoi -> la con cua dong 1.
    example_rows = [
        ("Ví dụ: Giấy in/photocopy A4 (dòng mẹ)", "Ram", 2, 70000, ""),
        ("Ví dụ: Giấy in khổ A3 (dòng CON của dòng 1 — đánh dấu x ở cột cuối)", "Ram", 1, 90000, "x"),
        ("Ví dụ: Mực in (dòng độc lập, không phải dòng con)", "Hộp", 1, 350000, ""),
    ]
    start_row = 3
    for offset, (desc, unit, qty, price, is_sub) in enumerate(example_rows):
        r = start_row + offset
        ws.cell(row=r, column=1, value=desc).alignment = st["left"]
        ws.cell(row=r, column=2, value=unit).alignment = st["center"]
        ws.cell(row=r, column=3, value=qty).alignment = st["center"]
        ws.cell(row=r, column=4, value=price).alignment = st["center"]
        ws.cell(row=r, column=5, value=is_sub).alignment = st["center"]
        for c in range(1, 6):
            ws.cell(row=r, column=c).font = st["normal"]
            ws.cell(row=r, column=c).border = st["border"]

    last_example_row = start_row + len(example_rows) - 1
    blank_rows_end = last_example_row + 30
    for r in range(last_example_row + 1, blank_rows_end + 1):
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = st["border"]
            ws.cell(row=r, column=c).font = st["normal"]

    dv_sub = DataValidation(type="list", formula1='"x"', allow_blank=True, showDropDown=False)
    dv_sub.error = "Chi dien 'x' neu la dong con, hoac de trong neu la dong doc lap."
    dv_sub.errorTitle = "Gia tri khong hop le"
    ws.add_data_validation(dv_sub)
    dv_sub.add(f"E{start_row}:E{blank_rows_end}")

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 22
    ws.freeze_panes = "A3"

    _write_instructions_sheet(wb, "HƯỚNG DẪN ĐIỀN FILE MẪU — Danh sách chi phí khác", [
        "",
        "1. Mỗi dòng là 1 khoản chi phí khác (vật tư, vật liệu, chi phí phụ...).",
        "2. Cột 'Khối lượng' và 'Đơn giá' là số, không điền chữ.",
        "3. TRƯỜNG HỢP CÓ HẠNG MỤC MẸ — HẠNG MỤC CON (dòng con nằm thụt vào): điền cột 'Là dòng con?",
        "   (x = có)' = 'x' cho những dòng là CON của dòng NGAY PHÍA TRÊN nó. Để TRỐNG nếu là dòng độc lập/",
        "   dòng mẹ. Nhiều dòng liên tiếp đánh dấu 'x' đều là con của cùng 1 dòng mẹ gần nhất phía trên.",
        "   Dòng ĐẦU TIÊN của cả file KHÔNG được đánh dấu 'x'. Xem dòng ví dụ số 2 (dòng con) ở trên.",
        "4. Sau khi điền xong, lưu file lại (giữ nguyên định dạng .xlsx), vào phần mềm bấm nút",
        "   'Nhập từ Excel' ở phần 'Chi phí khác' của hạng mục cần nhập, chọn đúng file này.",
        "5. Import sẽ THAY THẾ toàn bộ danh sách chi phí khác hiện tại của hạng mục đó — anh nên kiểm tra",
        "   lại sau khi import.",
    ])

    return wb


def parse_other_costs_excel(path):
    """Tra ve (rows, warnings) voi rows la danh sach otherRows:
    [{desc, unit, quantity, unitPrice, isSub}, ...]. Cot 'La dong con?' (cot E) danh dau 'x'
    -> dong do la CON cua dong ngay phia truoc no; dong dau tien cua file khong duoc la dong con."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Du lieu"] if "Du lieu" in wb.sheetnames else wb.worksheets[0]

    header_row = _find_header_row(ws, "don gia") or _find_header_row(ws, "mo ta")
    if header_row is None:
        raise ValueError("Không tìm thấy dòng tiêu đề trong file — anh dùng đúng file mẫu nhé.")

    rows = []
    warnings = []
    for r in range(header_row + 1, ws.max_row + 1):
        desc = ws.cell(row=r, column=1).value
        unit = ws.cell(row=r, column=2).value
        qty = ws.cell(row=r, column=3).value
        price = ws.cell(row=r, column=4).value
        sub_flag = ws.cell(row=r, column=5).value
        if all(v is None or str(v).strip() == "" for v in [desc, unit, qty, price, sub_flag]):
            continue
        desc_text = str(desc).strip() if desc is not None else ""
        if not desc_text:
            warnings.append(f"Dòng {r}: bỏ qua vì thiếu cột 'Mô tả'.")
            continue
        try:
            qty_val = float(qty) if qty is not None and str(qty).strip() != "" else 0
        except (TypeError, ValueError):
            qty_val = 0
        try:
            price_val = float(price) if price is not None and str(price).strip() != "" else 0
        except (TypeError, ValueError):
            price_val = 0
        is_sub = _is_sub_flag_true(sub_flag)
        if is_sub and not rows:
            warnings.append(f"Dòng {r}: đánh dấu là dòng con nhưng đây là khoản chi phí ĐẦU TIÊN của file (chưa có dòng mẹ phía trên) — đã tự bỏ đánh dấu, chuyển thành dòng độc lập.")
            is_sub = False
        rows.append({
            "desc": desc_text,
            "unit": (str(unit).strip() if unit is not None else ""),
            "quantity": qty_val,
            "unitPrice": price_val,
            "isSub": is_sub,
        })

    if not rows:
        raise ValueError("Không đọc được dòng dữ liệu nào hợp lệ trong file.")

    return rows, warnings


# ---- 3. Bang don gia tien luong chuyen gia (salaryTable) ----

def build_salary_template(salary_table):
    """Tao file mau da DIEN SAN theo du lieu Bang luong chuyen gia hien tai cua bao gia,
    de nguoi dung chi can sua so lieu roi import lai (thay vi go tu dau)."""
    st = _template_base_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Du lieu"

    ws.merge_cells("A1:C1")
    title_cell = ws.cell(
        row=1, column=1,
        value="FILE MAU — Bảng đơn giá tiền lương chuyên gia (đã điền sẵn dữ liệu hiện tại — chỉnh sửa rồi Import lại)",
    )
    title_cell.font = Font(bold=True, size=12, name="Times New Roman", color="8A3B2D")
    title_cell.alignment = st["left"]

    ws.cell(row=2, column=1, value="Thông tư áp dụng:").font = st["bold"]
    ws.cell(row=2, column=2, value=salary_table.get("circularName", "")).font = st["normal"]
    ws.cell(row=3, column=1, value="Số ngày công quy đổi/tháng:").font = st["bold"]
    ws.cell(row=3, column=2, value=salary_table.get("workDaysPerMonth", 26)).font = st["normal"]

    header_row = 5
    headers = ["Tên mức", "Định mức/tháng (VNĐ)", "Diễn giải (theo thông tư)"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = st["bold"]
        c.fill = st["header_fill"]
        c.alignment = st["center"]
        c.border = st["border"]

    levels = salary_table.get("levels", [])
    for offset, lvl in enumerate(levels):
        r = header_row + 1 + offset
        ws.cell(row=r, column=1, value=lvl.get("name", "")).alignment = st["left"]
        ws.cell(row=r, column=2, value=lvl.get("monthlyRate", 0)).alignment = st["center"]
        ws.cell(row=r, column=3, value=lvl.get("explanation", "")).alignment = st["left"]
        for c in range(1, 4):
            ws.cell(row=r, column=c).font = st["normal"]
            ws.cell(row=r, column=c).border = st["border"]

    blank_start = header_row + 1 + len(levels)
    for r in range(blank_start, blank_start + 10):
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = st["border"]
            ws.cell(row=r, column=c).font = st["normal"]

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 60
    ws.freeze_panes = f"A{header_row + 1}"

    _write_instructions_sheet(wb, "HƯỚNG DẪN ĐIỀN FILE MẪU — Bảng đơn giá tiền lương chuyên gia", [
        "",
        "1. File này đã được điền sẵn theo dữ liệu Bảng lương chuyên gia hiện tại của báo giá — anh chỉ cần sửa số liệu.",
        "2. 'Tên mức' ở dòng nào TRÙNG với tên mức đang có sẽ được CẬP NHẬT số liệu (giữ nguyên liên kết với các",
        "   hạng mục đang dùng mức đó). Tên mức MỚI (chưa có) sẽ được THÊM VÀO. Mức nào đang có nhưng KHÔNG còn",
        "   trong file sẽ bị XÓA — các hạng mục đang dùng mức bị xóa sẽ được phần mềm tự chuyển sang mức đầu tiên",
        "   còn lại, anh nên kiểm tra lại sau khi import.",
        "3. Có thể sửa ở dòng 2-3 (Thông tư áp dụng / Số ngày công quy đổi/tháng) nếu cần.",
        "4. Sau khi điền xong, lưu file lại (giữ nguyên định dạng .xlsx), vào phần mềm bấm nút",
        "   'Nhập từ Excel' ở phần 'Bảng lương chuyên gia', chọn đúng file này.",
    ])

    return wb


def parse_salary_excel(path):
    """Tra ve (data, warnings) voi data = {levels, workDaysPerMonth, circularName}."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Du lieu"] if "Du lieu" in wb.sheetnames else wb.worksheets[0]

    circular_name = ""
    work_days = 26
    for r in range(1, min(6, ws.max_row) + 1):
        label = ws.cell(row=r, column=1).value
        val = ws.cell(row=r, column=2).value
        if not label:
            continue
        label_norm = _strip_diacritics(label).strip().lower()
        if "thong tu" in label_norm:
            circular_name = str(val).strip() if val is not None else ""
        elif "ngay cong" in label_norm:
            try:
                work_days = float(val) if val is not None and str(val).strip() != "" else 26
            except (TypeError, ValueError):
                work_days = 26

    header_row = _find_header_row(ws, "dinh muc") or _find_header_row(ws, "định mức") or _find_header_row(ws, "ten muc") or _find_header_row(ws, "tên mức")
    if header_row is None:
        raise ValueError("Không tìm thấy dòng tiêu đề trong file — anh dùng đúng file mẫu nhé.")

    levels = []
    warnings = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        rate = ws.cell(row=r, column=2).value
        explanation = ws.cell(row=r, column=3).value
        if all(v is None or str(v).strip() == "" for v in [name, rate, explanation]):
            continue
        name_text = str(name).strip() if name is not None else ""
        if not name_text:
            warnings.append(f"Dòng {r}: bỏ qua vì thiếu 'Tên mức'.")
            continue
        try:
            rate_val = float(rate) if rate is not None and str(rate).strip() != "" else 0
        except (TypeError, ValueError):
            rate_val = 0
        levels.append({
            "name": name_text,
            "monthlyRate": rate_val,
            "explanation": (str(explanation).strip() if explanation is not None else ""),
        })

    if not levels:
        raise ValueError("Không đọc được mức chuyên gia nào hợp lệ trong file.")

    return {"levels": levels, "workDaysPerMonth": work_days, "circularName": circular_name}, warnings
